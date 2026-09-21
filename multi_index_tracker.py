#!/usr/bin/env python3
"""Build multi-index NIFTY valuation/backtest data for GitHub Pages.

This script is additive to the existing NIFTY 50 tracker. It uses the official
NSE daily multi-index archive file (ind_close_all_DDMMYYYY.csv), so one network
request gives close and valuation fields for many NIFTY indices at once.

Outputs under docs/data/:
  multi_index_catalog.json
  multi_index_latest.json
  multi_index_backtests.json
  multi_index_history.csv

Persistent canonical data at repository root:
  multi_index_quarterly.csv
  multi_index_monthly.csv

The first run bootstraps:
  * quarter-end snapshots from Mar-2012 onward (for forward-return backtests)
  * month-end snapshots from Apr-2021 onward (for current consolidated-era PE percentile)
Subsequent runs append only newly completed periods.
"""
from __future__ import annotations

import argparse
import html
import io
import json
import math
import os
import re
import smtplib
import time
from calendar import monthrange
from dataclasses import dataclass
from datetime import date, datetime, timedelta
from email.message import EmailMessage
from pathlib import Path
from statistics import median
from typing import Dict, Iterable, List, Optional, Tuple
from urllib.parse import urlencode

import pandas as pd
import requests
from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

NSE_ARCHIVE_BASE = "https://nsearchives.nseindia.com/content/indices"
NSE_SOURCE_URL = "https://www.nseindia.com/all-reports"
UA = "Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 Chrome/140 Safari/537.36"
MODEL_VERSION = "multi-index-1.0"
REGIME_SEAM = pd.Timestamp("2021-03-31")
MIN_EXPANDING_OBS = 8
MIN_LIVE_MONTHS = 8
MIN_BACKTEST_QUARTERS = 32
MIN_MATURE_3Y_OBS = 16
BOOTSTRAP_QUARTER_START = date(2012, 3, 31)
BOOTSTRAP_MONTH_START = date(2021, 4, 30)

# Official broad-market and sectoral families from NSE pages. Normalised names are
# used only for presentation grouping; any other PE-bearing equity index is kept
# under Strategy / Thematic / Other.
BROAD_NAMES = {
    "NIFTY50","NIFTYNEXT50","NIFTY100","NIFTYNEXT100","NIFTY200",
    "NIFTYTOTALMARKET","NIFTY500","NIFTY500MULTICAP502525",
    "NIFTY500LARGEMIDSMALLEQUALCAPWEIGHTED","NIFTYMIDCAP150","NIFTYMIDCAP50",
    "NIFTYMIDCAPSELECT","NIFTYMIDCAP100","NIFTYSMALLCAP500","NIFTYSMALLCAP250",
    "NIFTYSMALLCAP50","NIFTYSMALLCAP100","NIFTYMICROCAP250","NIFTYLARGEMIDCAP250",
    "NIFTYMIDSMALLCAP400","NIFTYMIDSMALLCAP4005050","NIFTYINDIAFPI150"
}
SECTOR_NAMES = {
    "NIFTYAUTO","NIFTYBANK","NIFTYCEMENT","NIFTYCAPITALGOODS","NIFTYCHEMICALS",
    "NIFTYCOMMERCIALTRANSPORTSERVICES","NIFTYCONSTRUCTION","NIFTYCONSUMERSERVICES",
    "NIFTYFINANCIALSERVICES","NIFTYFINANCIALSERVICES2550","NIFTYFINANCIALSERVICESEXBANK",
    "NIFTYFMCG","NIFTYHEALTHCARE","NIFTYHEALTHCAREINDEX","NIFTYHOSPITALS",
    "NIFTYHOUSINGFINANCE","NIFTYINSURANCE","NIFTYIT","NIFTYMEDIA","NIFTYMETAL",
    "NIFTYNBFC","NIFTYPHARMA","NIFTYPOWER","NIFTYPRIVATEBANK","NIFTYPSUBANK",
    "NIFTYREALTY","NIFTYREITSREALTY","NIFTYRETAIL","NIFTYTELECOMMUNICATIONS",
    "NIFTYCONSUMERDURABLES","NIFTYOILGAS","NIFTY500HEALTHCARE",
    "NIFTYMIDSMALLFINANCIALSERVICES","NIFTYMIDSMALLHEALTHCARE","NIFTYMIDSMALLITTELECOM"
}

# Common archive abbreviations -> polished display label.
DISPLAY_MAP = {
    "NIFTY50":"NIFTY 50",
    "NIFTYNEXT50":"NIFTY Next 50",
    "NIFTY100":"NIFTY 100",
    "NIFTY200":"NIFTY 200",
    "NIFTY500":"NIFTY 500",
    "NIFTYTOTALMKT":"NIFTY Total Market",
    "NIFTYTOTALMARKET":"NIFTY Total Market",
    "NIFTYMIDCAP150":"NIFTY Midcap 150",
    "NIFTYMIDCAP50":"NIFTY Midcap 50",
    "NIFTYMIDCAP100":"NIFTY Midcap 100",
    "NIFTYMIDSELECT":"NIFTY Midcap Select",
    "NIFTYMIDCAPSELECT":"NIFTY Midcap Select",
    "NIFTYSMLCAP250":"NIFTY Smallcap 250",
    "NIFTYSMALLCAP250":"NIFTY Smallcap 250",
    "NIFTYSMLCAP100":"NIFTY Smallcap 100",
    "NIFTYSMALLCAP100":"NIFTY Smallcap 100",
    "NIFTYSMLCAP50":"NIFTY Smallcap 50",
    "NIFTYSMALLCAP50":"NIFTY Smallcap 50",
    "NIFTYMICROCAP250":"NIFTY Microcap 250",
    "NIFTYLARGEMID250":"NIFTY LargeMidcap 250",
    "NIFTYLARGEMIDCAP250":"NIFTY LargeMidcap 250",
    "NIFTYBANK":"NIFTY Bank",
    "NIFTYFINSERVICE":"NIFTY Financial Services",
    "NIFTYFINANCIALSERVICES":"NIFTY Financial Services",
    "NIFTYIT":"NIFTY IT",
    "NIFTYAUTO":"NIFTY Auto",
    "NIFTYFMCG":"NIFTY FMCG",
    "NIFTYPHARMA":"NIFTY Pharma",
    "NIFTYHEALTHCAREINDEX":"NIFTY Healthcare",
    "NIFTYHEALTHCARE":"NIFTY Healthcare",
    "NIFTYMETAL":"NIFTY Metal",
    "NIFTYREALTY":"NIFTY Realty",
    "NIFTYENERGY":"NIFTY Energy",
    "NIFTYOILANDGAS":"NIFTY Oil & Gas",
    "NIFTYOILGAS":"NIFTY Oil & Gas",
    "NIFTYMEDIA":"NIFTY Media",
    "NIFTYPSUBANK":"NIFTY PSU Bank",
    "NIFTYPVTBANK":"NIFTY Private Bank",
    "NIFTYPRIVATEBANK":"NIFTY Private Bank",
    "NIFTYCONSRDURBL":"NIFTY Consumer Durables",
    "NIFTYCONSUMERDURABLES":"NIFTY Consumer Durables",
    "NIFTYFINANCIALSERVICES2550":"NIFTY Financial Services 25/50",
}

# Equity-like exclusions where PE-based valuation is not economically meaningful.
EXCLUDE_TERMS = (
    "GSEC", "G-SEC", " GS ", "BOND", "SDL", "DEBT", "LIQUID", "MONEY MARKET",
    "ARBITRAGE", "FUTURES", "DIVIDEND POINT", "INVERSE", "LEVERAGE",
    "EQUITY SAVINGS", "DYNAMIC P/E", "DYNAMIC PE", "DYNAMIC P/B", "DYNAMIC PB"
)


def norm_name(s: str) -> str:
    return re.sub(r"[^A-Z0-9]", "", str(s).upper())




def canonical_key(raw: str) -> str:
    n = norm_name(raw)
    return norm_name(DISPLAY_MAP[n]) if n in DISPLAY_MAP else n

def slugify(s: str) -> str:
    s = re.sub(r"[^a-z0-9]+", "-", str(s).lower()).strip("-")
    return s or "index"


def prettify(raw: str) -> str:
    n = norm_name(raw)
    if n in DISPLAY_MAP:
        return DISPLAY_MAP[n]
    # Conservative title casing with NIFTY kept uppercase.
    x = re.sub(r"\s+", " ", str(raw).strip())
    words = []
    for w in x.split(" "):
        u = w.upper()
        if u in {"NIFTY","IT","FMCG","PSU","NBFC","REIT","REITS","ESG","FPI"}:
            words.append(u)
        elif u in {"50","100","150","200","250","400","500"}:
            words.append(u)
        else:
            words.append(w.title())
    return " ".join(words)


def classify_group(name: str) -> str:
    n = norm_name(name)
    # Handle common archive abbreviations before official-list comparison.
    alias = {
        "NIFTYTOTALMKT":"NIFTYTOTALMARKET",
        "NIFTYMIDSELECT":"NIFTYMIDCAPSELECT",
        "NIFTYSMLCAP250":"NIFTYSMALLCAP250",
        "NIFTYSMLCAP100":"NIFTYSMALLCAP100",
        "NIFTYSMLCAP50":"NIFTYSMALLCAP50",
        "NIFTYLARGEMID250":"NIFTYLARGEMIDCAP250",
        "NIFTYFINSERVICE":"NIFTYFINANCIALSERVICES",
        "NIFTYPVTBANK":"NIFTYPRIVATEBANK",
        "NIFTYCONSRDURBL":"NIFTYCONSUMERDURABLES",
        "NIFTYOILANDGAS":"NIFTYOILGAS",
    }.get(n, n)
    if alias in BROAD_NAMES:
        return "Broad Market"
    if alias in SECTOR_NAMES or n == "NIFTYENERGY":
        return "Sectoral"
    return "Strategy / Thematic / Other"


def is_equity_like(name: str) -> bool:
    u = f" {str(name).upper()} "
    if not u.strip().startswith("NIFTY"):
        return False
    return not any(term in u for term in EXCLUDE_TERMS)


def _norm_col(c: str) -> str:
    return re.sub(r"[^a-z0-9]", "", str(c).strip().lower())


def _find_col(df: pd.DataFrame, candidates: Iterable[str]) -> Optional[str]:
    normalized = {_norm_col(c): c for c in df.columns}
    for cand in candidates:
        if cand in normalized:
            return normalized[cand]
    for norm, orig in normalized.items():
        if any(cand in norm for cand in candidates):
            return orig
    return None


def session() -> requests.Session:
    s = requests.Session()
    s.headers.update({
        "User-Agent": UA,
        "Accept-Language": "en-US,en;q=0.9",
        "Accept": "text/csv,text/plain,*/*",
        "Referer": "https://www.nseindia.com/",
    })
    return s


def fetch_snapshot_exact(s: requests.Session, d: date) -> Tuple[date, pd.DataFrame, str]:
    url = f"{NSE_ARCHIVE_BASE}/ind_close_all_{d:%d%m%Y}.csv"
    r = s.get(url, timeout=30)
    if r.status_code == 404:
        raise FileNotFoundError(url)
    r.raise_for_status()
    body = r.text.lstrip("\ufeff \t\r\n")
    if not body or body.startswith("<"):
        raise ValueError(f"Non-CSV response for {d}")
    df = pd.read_csv(io.StringIO(r.text))

    name_col = _find_col(df, ["indexname", "index"])
    close_col = _find_col(df, ["closingindexvalue", "closing", "close", "closingvalue"])
    pe_col = _find_col(df, ["pe", "peratio", "priceearnings", "priceearningsratio"])
    date_col = _find_col(df, ["indexdate", "date"])
    if name_col is None or close_col is None:
        raise ValueError(f"Could not identify index/close columns: {list(df.columns)}")

    out = pd.DataFrame({
        "IndexNameRaw": df[name_col].astype(str).str.strip(),
        "Close": pd.to_numeric(df[close_col], errors="coerce"),
        "PE": pd.to_numeric(df[pe_col], errors="coerce") if pe_col else float("nan"),
    })
    out = out[out["Close"].notna() & (out["Close"] > 0)].copy()
    out = out[out["IndexNameRaw"].map(is_equity_like)].copy()
    out["IndexKey"] = out["IndexNameRaw"].map(canonical_key)
    out["IndexName"] = out["IndexNameRaw"].map(prettify)
    out["Group"] = out["IndexNameRaw"].map(classify_group)

    out_date = d
    if date_col is not None:
        vals = pd.to_datetime(df[date_col], errors="coerce", dayfirst=True).dropna()
        if len(vals):
            out_date = vals.iloc[0].date()
    return out_date, out, url


def fetch_snapshot_on_or_before(s: requests.Session, target: date, max_lookback: int = 12) -> Tuple[date, pd.DataFrame, str]:
    errors = []
    for i in range(max_lookback + 1):
        d = target - timedelta(days=i)
        try:
            return fetch_snapshot_exact(s, d)
        except FileNotFoundError:
            continue
        except Exception as e:
            errors.append(f"{d}: {e}")
            continue
    tail = "; ".join(errors[-3:]) if errors else "no trading-day archive found"
    raise RuntimeError(f"Unable to fetch NSE multi-index snapshot through {target}: {tail}")


def quarter_ends(start: date, end: date) -> List[date]:
    out = []
    y, m = start.year, start.month
    # snap to quarter-end month
    qm = ((m - 1) // 3 + 1) * 3
    cur = date(y, qm, monthrange(y, qm)[1])
    if cur < start:
        qm += 3
        if qm > 12:
            qm -= 12; y += 1
        cur = date(y, qm, monthrange(y, qm)[1])
    while cur <= end:
        out.append(cur)
        qm = cur.month + 3
        y = cur.year
        if qm > 12:
            qm -= 12; y += 1
        cur = date(y, qm, monthrange(y, qm)[1])
    return out


def month_ends(start: date, end: date) -> List[date]:
    out = []
    y, m = start.year, start.month
    cur = date(y, m, monthrange(y, m)[1])
    if cur < start:
        m += 1
        if m > 12:
            m = 1; y += 1
        cur = date(y, m, monthrange(y, m)[1])
    while cur <= end:
        out.append(cur)
        m = cur.month + 1
        y = cur.year
        if m > 12:
            m = 1; y += 1
        cur = date(y, m, monthrange(y, m)[1])
    return out


def latest_completed_quarter_end(d: date) -> date:
    qm = ((d.month - 1)//3 + 1)*3
    qe = date(d.year, qm, monthrange(d.year, qm)[1])
    if d >= qe:
        return qe
    qm -= 3
    y = d.year
    if qm <= 0:
        qm += 12; y -= 1
    return date(y, qm, monthrange(y, qm)[1])


def latest_completed_month_end(d: date) -> date:
    me = date(d.year, d.month, monthrange(d.year, d.month)[1])
    if d >= me:
        return me
    m = d.month - 1; y = d.year
    if m == 0:
        m = 12; y -= 1
    return date(y, m, monthrange(y, m)[1])


def append_period_snapshots(
    s: requests.Session,
    path: Path,
    periods: List[date],
    period_kind: str,
    sleep_seconds: float = 0.10,
) -> pd.DataFrame:
    cols = ["IndexKey","IndexName","Group","Date","SourceDate","Close","PE"]
    if path.exists():
        df = pd.read_csv(path)
    else:
        df = pd.DataFrame(columns=cols)

    have = set()
    if len(df):
        dates = pd.to_datetime(df["Date"], errors="coerce")
        have = {x.date() for x in dates.dropna().unique()}

    missing = [p for p in periods if p not in have]
    if not missing:
        return df

    additions = []
    for i, period_end in enumerate(missing, 1):
        try:
            source_date, snap, _ = fetch_snapshot_on_or_before(s, period_end, max_lookback=12)
        except Exception as e:
            print(f"WARNING: {period_kind} bootstrap skipped {period_end:%d-%b-%Y}: {e}")
            continue
        for _, r in snap.iterrows():
            additions.append({
                "IndexKey": r["IndexKey"],
                "IndexName": r["IndexName"],
                "Group": r["Group"],
                "Date": period_end.isoformat(),
                "SourceDate": source_date.isoformat(),
                "Close": float(r["Close"]),
                "PE": None if pd.isna(r["PE"]) else float(r["PE"]),
            })
        print(f"{period_kind}: added {period_end:%d-%b-%Y} from NSE {source_date:%d-%b-%Y} ({len(snap)} indices)")
        if sleep_seconds:
            time.sleep(sleep_seconds)

    out = pd.concat([df, pd.DataFrame(additions)], ignore_index=True, sort=False)
    out["Date"] = pd.to_datetime(out["Date"], errors="coerce")
    out["SourceDate"] = pd.to_datetime(out["SourceDate"], errors="coerce")
    out["Close"] = pd.to_numeric(out["Close"], errors="coerce")
    out["PE"] = pd.to_numeric(out["PE"], errors="coerce")
    out = out.dropna(subset=["IndexKey","Date","Close"])
    out = out.sort_values(["IndexKey","Date","SourceDate"]).drop_duplicates(["IndexKey","Date"], keep="last")
    path.parent.mkdir(parents=True, exist_ok=True)
    out.to_csv(path, index=False, date_format="%Y-%m-%d", float_format="%.12g")
    return out


def percentile_midrank(values: Iterable[float], x: float) -> float:
    vals = [float(v) for v in values if pd.notna(v)]
    if not vals:
        return float("nan")
    lt = sum(v < x for v in vals)
    eq = sum(abs(v - x) < 1e-10 for v in vals)
    return (lt + 0.5*eq) / len(vals)


def rank_pct(values: Iterable[float], x: float) -> float:
    vals = [float(v) for v in values if pd.notna(v)]
    if not vals:
        return float("nan")
    lt = sum(v < x for v in vals)
    eq = sum(abs(v-x) < 1e-10 for v in vals)
    return (lt + (eq + 1)/2.0) / len(vals)


def growth_score(g: float) -> float:
    if pd.isna(g): return float("nan")
    if g <= 0: return 20.0
    if g < 0.05: return 20 + 400*g
    if g < 0.10: return 40 + 400*(g-0.05)
    if g < 0.15: return 60 + 400*(g-0.10)
    if g < 0.25: return 80 + 200*(g-0.15)
    return 100.0


def valuation_quintile(p: float) -> str:
    if pd.isna(p): return "Insufficient history"
    if p <= .20: return "Q1 Cheapest"
    if p <= .40: return "Q2"
    if p <= .60: return "Q3"
    if p <= .80: return "Q4"
    return "Q5 Most Expensive"


def linear_quantile(values: Iterable[float], q: float) -> float:
    vals = sorted(float(v) for v in values if pd.notna(v))
    if not vals:
        return float("nan")
    q = max(0.0, min(1.0, float(q)))
    if len(vals) == 1:
        return vals[0]
    pos = (len(vals)-1)*q
    lo = int(math.floor(pos)); hi = int(math.ceil(pos))
    if lo == hi: return vals[lo]
    w = pos-lo
    return vals[lo]*(1-w)+vals[hi]*w


def signal_boundaries(pe_values: List[float], growth_score_value: float, current_eps: float, current_close: float):
    if len(pe_values) < MIN_LIVE_MONTHS or pd.isna(growth_score_value):
        return None
    wv, wg = .70, .30
    def one(threshold):
        req_val = (threshold - wg*growth_score_value)/wv
        req_val = max(0.0, min(100.0, req_val))
        target_pct = 1 - req_val/100.0
        pe = linear_quantile(pe_values, target_pct)
        nifty = current_eps*pe
        move = nifty/current_close - 1 if current_close else float("nan")
        return {
            "score_threshold": threshold,
            "valuation_score_required": req_val,
            "pe_percentile": target_pct,
            "pe": pe,
            "nifty_level": nifty,
            "move_from_current": move,
        }
    return {"buy_hold": one(70), "hold_sell": one(40)}


def build_quarterly_backtest(qdf: pd.DataFrame, key: str) -> Tuple[pd.DataFrame, Dict]:
    d = qdf[qdf["IndexKey"] == key].copy().sort_values("Date")
    if d.empty:
        return d, {"rows": [], "meta": {"status":"no_history"}}
    d["Date"] = pd.to_datetime(d["Date"])
    d["PE"] = pd.to_numeric(d["PE"], errors="coerce")
    d["Close"] = pd.to_numeric(d["Close"], errors="coerce")
    d["Regime"] = d["Date"].apply(lambda x: "Consolidated" if x >= REGIME_SEAM else "Standalone")
    d["EPS"] = d["Close"] / d["PE"]

    # YoY implied EPS growth (quarter aligned, same PE regime)
    eps_map = dict(zip(d["Date"], d["EPS"]))
    reg_map = dict(zip(d["Date"], d["Regime"]))
    growth = []
    for _, r in d.iterrows():
        prior = r["Date"] - pd.DateOffset(years=1)
        if prior in eps_map and reg_map.get(prior) == r["Regime"] and pd.notna(eps_map[prior]) and pd.notna(r["EPS"]):
            growth.append(float(r["EPS"] / eps_map[prior] - 1))
        else:
            growth.append(float("nan"))
    d["EPS_Growth_YoY"] = growth

    # Descriptive regime percentile/quintile only on valid PE observations.
    reg_pct = []
    for _, r in d.iterrows():
        if pd.isna(r["PE"]):
            reg_pct.append(float("nan")); continue
        vals = d.loc[(d["Regime"] == r["Regime"]) & d["PE"].notna(), "PE"].tolist()
        reg_pct.append(rank_pct(vals, float(r["PE"])))
    d["PE_Pct_Regime"] = reg_pct
    d["Valuation_Quintile"] = d["PE_Pct_Regime"].map(valuation_quintile)

    # Forward returns use price observations whether or not future PE is available.
    price_map = dict(zip(d["Date"], d["Close"]))
    for years, col in [(1,"Fwd_1Y"),(3,"Fwd_3Y"),(5,"Fwd_5Y"),(10,"Fwd_10Y")]:
        vals = []
        for _, r in d.iterrows():
            future = r["Date"] + pd.DateOffset(years=years)
            if future in price_map and pd.notna(price_map[future]) and pd.notna(r["Close"]):
                vals.append(float((price_map[future]/r["Close"])**(1/years)-1))
            else:
                vals.append(float("nan"))
        d[col] = vals

    rows = []
    order = ["Q1 Cheapest","Q2","Q3","Q4","Q5 Most Expensive"]
    for q in order:
        x = d[d["Valuation_Quintile"] == q]
        def med(col):
            s = pd.to_numeric(x[col], errors="coerce").dropna()
            return None if s.empty else float(s.median())
        s3 = pd.to_numeric(x["Fwd_3Y"], errors="coerce").dropna()
        rows.append({
            "quintile": q,
            "n": int(len(x)),
            "median_1y": med("Fwd_1Y"),
            "median_3y": med("Fwd_3Y"),
            "median_5y": med("Fwd_5Y"),
            "median_10y": med("Fwd_10Y"),
            "loss_3y": None if s3.empty else float((s3 < 0).mean()),
            "n_3y": int(len(s3)),
        })

    valid_pe = int(d["PE"].notna().sum())
    first = d["Date"].min().date().isoformat() if len(d) else None
    last = d["Date"].max().date().isoformat() if len(d) else None
    status = "ok" if valid_pe >= MIN_BACKTEST_QUARTERS else "insufficient_history"
    return d, {"rows": rows, "meta": {"status":status,"valid_pe_quarters":valid_pe,"first_quarter":first,"last_quarter":last}}



def build_legacy_nifty50_backtest(path: Path) -> Optional[Dict]:
    """Build the canonical NIFTY 50 backtest summary from the original long-history file.

    The multi-index NSE archive only offers a much shorter comparable history for
    NIFTY 50.  When the original tracker file is present, preserve its 1999-onward
    quarter-end backtest rather than replacing it with the generic multi-index sample.
    """
    if not path.exists():
        return None
    try:
        d = pd.read_csv(path)
    except Exception as e:
        print(f"WARNING: could not read canonical NIFTY 50 backtest {path}: {e}")
        return None

    required = {"Date", "PE", "Valuation_Quintile", "Fwd_1Y", "Fwd_3Y", "Fwd_5Y", "Fwd_10Y"}
    if not required.issubset(set(d.columns)):
        print(f"WARNING: canonical NIFTY 50 backtest missing columns: {sorted(required-set(d.columns))}")
        return None

    d["Date"] = pd.to_datetime(d["Date"], errors="coerce")
    d["PE"] = pd.to_numeric(d["PE"], errors="coerce")
    for c in ["Fwd_1Y", "Fwd_3Y", "Fwd_5Y", "Fwd_10Y"]:
        d[c] = pd.to_numeric(d[c], errors="coerce")

    rows = []
    order = ["Q1 Cheapest", "Q2", "Q3", "Q4", "Q5 Most Expensive"]
    for q in order:
        x = d[d["Valuation_Quintile"] == q]
        def med(col):
            vals = x[col].dropna()
            return None if vals.empty else float(vals.median())
        s3 = x["Fwd_3Y"].dropna()
        rows.append({
            "quintile": q,
            "n": int(len(x)),
            "median_1y": med("Fwd_1Y"),
            "median_3y": med("Fwd_3Y"),
            "median_5y": med("Fwd_5Y"),
            "median_10y": med("Fwd_10Y"),
            "loss_3y": None if s3.empty else float((s3 < 0).mean()),
            "n_3y": int(len(s3)),
        })

    valid_pe = int(d["PE"].notna().sum())
    valid_dates = d["Date"].dropna()
    first = valid_dates.min().date().isoformat() if len(valid_dates) else None
    last = valid_dates.max().date().isoformat() if len(valid_dates) else None
    return {
        "rows": rows,
        "meta": {
            "status": "ok" if valid_pe >= MIN_BACKTEST_QUARTERS else "insufficient_history",
            "valid_pe_quarters": valid_pe,
            "first_quarter": first,
            "last_quarter": last,
            "source": "canonical_nifty50_long_history",
        },
    }

def find_prior_row(prior_snap: pd.DataFrame, key: str) -> Optional[pd.Series]:
    x = prior_snap[prior_snap["IndexKey"] == key]
    if x.empty: return None
    return x.iloc[0]


def build_current_for_index(
    key: str,
    latest_row: pd.Series,
    latest_date: date,
    prior_snap: pd.DataFrame,
    monthly: pd.DataFrame,
    quarter_bt: pd.DataFrame,
) -> Dict:
    close = float(latest_row["Close"])
    pe = None if pd.isna(latest_row["PE"]) else float(latest_row["PE"])
    prior = find_prior_row(prior_snap, key)
    prior_close = None if prior is None else float(prior["Close"])
    prior_pe = None if prior is None or pd.isna(prior["PE"]) else float(prior["PE"])

    monthly_idx = monthly[(monthly["IndexKey"] == key) & monthly["PE"].notna()].copy()
    monthly_idx["Date"] = pd.to_datetime(monthly_idx["Date"], errors="coerce")
    monthly_idx = monthly_idx[monthly_idx["Date"] >= REGIME_SEAM]
    pe_hist = [float(v) for v in monthly_idx["PE"].dropna().tolist()]

    eps = close/pe if pe and pe > 0 else None
    prior_eps = prior_close/prior_pe if prior_close and prior_pe and prior_pe > 0 else None
    growth = (eps/prior_eps - 1) if eps and prior_eps else None
    gs = growth_score(growth) if growth is not None else None

    if pe is not None and len(pe_hist) >= MIN_LIVE_MONTHS:
        pct_val = percentile_midrank(pe_hist + [pe], pe)
        valuation_score = 100*(1-pct_val)
    else:
        pct_val = None
        valuation_score = None

    if valuation_score is not None and gs is not None:
        composite = .70*valuation_score + .30*gs
        signal = "BUY" if composite >= 70 else ("HOLD" if composite >= 40 else "SELL")
    else:
        composite = None
        signal = None

    bounds = signal_boundaries(pe_hist + ([pe] if pe is not None else []), gs, eps, close) if eps and gs is not None else None
    med_pe = median(pe_hist) if pe_hist else None
    q = valuation_quintile(pct_val) if pct_val is not None else "Insufficient history"

    return {
        "index_key": key,
        "index_name": latest_row["IndexName"],
        "group": latest_row["Group"],
        "slug": slugify(latest_row["IndexName"]),
        "as_of": latest_date.isoformat(),
        "nifty_close": close,
        "pe": pe,
        "earnings_yield": (1/pe if pe else None),
        "implied_eps": eps,
        "yoy_eps_growth": growth,
        "pe_percentile": pct_val,
        "era_median_pe": med_pe,
        "valuation_quintile": q,
        "valuation_score": valuation_score,
        "growth_score": gs,
        "composite_score": composite,
        "signal": signal,
        "signal_boundaries": bounds,
        "live_pe_history_months": len(pe_hist),
        "backtest_valid_pe_quarters": int(quarter_bt["PE"].notna().sum()) if not quarter_bt.empty else 0,
        "model_version": MODEL_VERSION,
        "source": "NSE Indices daily archive",
    }


def append_daily_history(path: Path, latest_map: Dict[str, Dict]):
    cols = ["date","slug","index_name","group","close","pe","pe_percentile","yoy_eps_growth","valuation_score","growth_score","composite_score","signal"]
    if path.exists():
        hist = pd.read_csv(path)
        # Prune previously tracked indices that are no longer useful enough to
        # appear on the dashboard. Raw monthly/quarterly source history is kept
        # separately so an index can automatically re-enter later once it has
        # sufficient history.
        hist = hist[hist["slug"].isin(set(latest_map.keys()))].copy()
    else:
        hist = pd.DataFrame(columns=cols)
    rows = []
    for slug, x in latest_map.items():
        rows.append({
            "date": x["as_of"], "slug": slug, "index_name": x["index_name"], "group": x["group"],
            "close": x["nifty_close"], "pe": x["pe"], "pe_percentile": x["pe_percentile"],
            "yoy_eps_growth": x["yoy_eps_growth"], "valuation_score": x["valuation_score"],
            "growth_score": x["growth_score"], "composite_score": x["composite_score"], "signal": x["signal"],
        })
    out = pd.concat([hist, pd.DataFrame(rows)], ignore_index=True, sort=False)
    out = out.drop_duplicates(["date","slug"], keep="last").sort_values(["date","slug"])
    path.parent.mkdir(parents=True, exist_ok=True)
    out.to_csv(path, index=False, float_format="%.12g")


def make_catalog(latest_map: Dict[str, Dict], backtests: Dict[str, Dict]) -> Dict:
    items = []
    for slug, x in latest_map.items():
        bt_meta = backtests.get(slug, {}).get("meta", {})
        if x["pe"] is None:
            status = "pe_unavailable"
        elif x["composite_score"] is None:
            status = "insufficient_history"
        else:
            status = "ok"
        items.append({
            "slug": slug,
            "name": x["index_name"],
            "group": x["group"],
            "status": status,
            "valid_pe_quarters": bt_meta.get("valid_pe_quarters", 0),
            "matured_3y_obs": sum(int(r.get("n_3y", 0) or 0) for r in backtests.get(slug, {}).get("rows", [])),
            "live_pe_history_months": x.get("live_pe_history_months", 0),
        })
    group_order = {"Broad Market":0,"Sectoral":1,"Strategy / Thematic / Other":2}
    items.sort(key=lambda z:(group_order.get(z["group"],9), z["name"]))
    return {
        "generated_at": datetime.utcnow().isoformat(timespec="seconds") + "Z",
        "default_slug": next((i["slug"] for i in items if i["name"] == "NIFTY 50"), items[0]["slug"] if items else None),
        "groups": ["Broad Market","Sectoral","Strategy / Thematic / Other"],
        "items": items,
    }



def load_legacy_nifty50_detail(path: Path) -> Optional[pd.DataFrame]:
    """Return the canonical NIFTY 50 quarter-end detail for workbook export."""
    if not path.exists():
        return None
    try:
        d = pd.read_csv(path)
    except Exception as e:
        print(f"WARNING: could not read NIFTY 50 workbook-detail history {path}: {e}")
        return None
    if "Date" not in d.columns or "PE" not in d.columns:
        return None
    d["Date"] = pd.to_datetime(d["Date"], errors="coerce")
    d = d[d["Date"].notna()].copy().sort_values("Date")
    if "Price" in d.columns and "Close" not in d.columns:
        d["Close"] = pd.to_numeric(d["Price"], errors="coerce")
    d["PE"] = pd.to_numeric(d["PE"], errors="coerce")
    if "Regime" not in d.columns:
        d["Regime"] = d["Date"].apply(lambda x: "Consolidated" if x >= REGIME_SEAM else "Standalone")
    return d


def _safe_excel_value(v):
    if v is None:
        return None
    if isinstance(v, float) and (math.isnan(v) or math.isinf(v)):
        return None
    if pd.isna(v):
        return None
    if isinstance(v, pd.Timestamp):
        return v.to_pydatetime()
    return v


def _autosize(ws, min_width: int = 10, max_width: int = 34):
    for col in range(1, ws.max_column + 1):
        letter = get_column_letter(col)
        width = 0
        for row in range(1, min(ws.max_row, 250) + 1):
            v = ws.cell(row, col).value
            if v is None:
                continue
            width = max(width, len(str(v)))
        ws.column_dimensions[letter].width = min(max(width + 2, min_width), max_width)


def _style_workbook_sheet(ws, freeze: Optional[str] = None):
    ws.sheet_view.showGridLines = False
    if freeze:
        ws.freeze_panes = freeze


def _apply_table_header(ws, row: int, start_col: int, end_col: int):
    fill = PatternFill("solid", fgColor="0F2742")
    font = Font(color="FFFFFF", bold=True)
    for c in range(start_col, end_col + 1):
        cell = ws.cell(row, c)
        cell.fill = fill
        cell.font = font
        cell.alignment = Alignment(horizontal="center", vertical="center")


def _apply_section_header(ws, row: int, title: str, end_col: int = 6):
    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=end_col)
    c = ws.cell(row, 1, title)
    c.fill = PatternFill("solid", fgColor="0B1F36")
    c.font = Font(color="FFFFFF", bold=True, size=11)
    c.alignment = Alignment(horizontal="left")


def _set_imported_header_comment(cell, source_url: str):
    cell.comment = Comment(f"Imported source: {source_url}", "Index Valuation Tracker")


def _write_index_workbook(
    out_path: Path,
    x: Dict,
    summary: Dict,
    monthly_detail: pd.DataFrame,
    quarterly_detail: pd.DataFrame,
    history_detail: pd.DataFrame,
):
    """Create a self-contained Excel download for one eligible index."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Overview"
    _style_workbook_sheet(ws)

    dark_fill = PatternFill("solid", fgColor="071A2D")
    teal_fill = PatternFill("solid", fgColor="D9F2EE")
    gray_fill = PatternFill("solid", fgColor="EEF2F6")
    white_font = Font(color="FFFFFF", bold=True)
    blue_font = Font(color="0000FF")
    green_font = Font(color="008000")
    gray_font = Font(color="666666")
    black_font = Font(color="000000")
    orange_fill = PatternFill("solid", fgColor="FFF2CC")
    thin_gray = Side(style="thin", color="D6DEE8")
    top_border = Border(top=Side(style="thin", color="7B8794"))

    ws.merge_cells("A1:F1")
    ws["A1"] = f"{x['index_name']} Valuation Tracker"
    ws["A1"].fill = dark_fill
    ws["A1"].font = Font(color="FFFFFF", bold=True, size=16)
    ws["A1"].alignment = Alignment(horizontal="left", vertical="center")
    ws.row_dimensions[1].height = 26

    ws["A2"] = "As of"
    ws["B2"] = datetime.fromisoformat(x["as_of"])
    ws["B2"].number_format = "dd-mmm-yyyy"
    ws["D2"] = "Group"
    ws["E2"] = x.get("group")
    for c in ("A2","D2"):
        ws[c].font = gray_font
    for c in ("B2","E2"):
        ws[c].font = green_font

    _apply_section_header(ws, 4, "Current valuation and score", 6)
    metrics = [
        ("Index level", x.get("nifty_close"), '#,##0.00;[Red](#,##0.00);-'),
        ("P/E", x.get("pe"), '0.00x;[Red](0.00x);-'),
        ("P/E percentile", x.get("pe_percentile"), '0.0%'),
        ("Implied EPS", x.get("implied_eps"), '#,##0.0;[Red](#,##0.0);-'),
        ("YoY implied EPS growth", x.get("yoy_eps_growth"), '0.0%'),
        ("Earnings yield", x.get("earnings_yield"), '0.00%'),
        ("Valuation score", x.get("valuation_score"), '0.0'),
        ("Growth score", x.get("growth_score"), '0.0'),
        ("Composite score", x.get("composite_score"), '0.0'),
        ("Signal", x.get("signal"), '@'),
    ]
    r = 5
    for i in range(0, len(metrics), 2):
        left = metrics[i]
        right = metrics[i+1] if i+1 < len(metrics) else None
        ws.cell(r,1,left[0]).font = gray_font
        ws.cell(r,2,_safe_excel_value(left[1])).font = green_font if not isinstance(left[1], str) else black_font
        ws.cell(r,2).number_format = left[2]
        if right:
            ws.cell(r,4,right[0]).font = gray_font
            ws.cell(r,5,_safe_excel_value(right[1])).font = green_font if not isinstance(right[1], str) else black_font
            ws.cell(r,5).number_format = right[2]
        r += 1

    _apply_section_header(ws, 11, "Signal thresholds", 6)
    ws["A12"] = "Boundary"
    ws["B12"] = "P/E"
    ws["C12"] = "Equivalent index level"
    ws["D12"] = "% move from current"
    _apply_table_header(ws, 12, 1, 4)
    bounds = x.get("signal_boundaries") or {}
    threshold_rows = [("BUY → HOLD", bounds.get("buy_hold")), ("HOLD → SELL", bounds.get("hold_sell"))]
    rr = 13
    for label, b in threshold_rows:
        ws.cell(rr,1,label)
        if b:
            ws.cell(rr,2,_safe_excel_value(b.get("pe"))).number_format = '0.00x'
            ws.cell(rr,3,_safe_excel_value(b.get("nifty_level"))).number_format = '#,##0'
            ws.cell(rr,4,_safe_excel_value(b.get("move_from_current"))).number_format = '0.0%'
        rr += 1

    _apply_section_header(ws, 16, "Forward returns by starting valuation", 7)
    headers = ["Starting valuation","Observations","1Y median","3Y median CAGR","5Y median CAGR","10Y median CAGR","3Y loss frequency"]
    for c,h in enumerate(headers,1): ws.cell(17,c,h)
    _apply_table_header(ws,17,1,7)
    for i, row in enumerate(summary.get("rows", []), start=18):
        vals = [
            row.get("quintile"), row.get("n"), row.get("median_1y"), row.get("median_3y"),
            row.get("median_5y"), row.get("median_10y"), row.get("loss_3y")
        ]
        for c,v in enumerate(vals,1):
            ws.cell(i,c,_safe_excel_value(v))
        for c in range(3,8): ws.cell(i,c).number_format='0.0%'

    meta = summary.get("meta", {})
    note_row = 24
    _apply_section_header(ws, note_row, "Coverage and methodology", 6)
    notes = [
        ("Monthly P/E observations", x.get("live_pe_history_months")),
        ("Quarter-end P/E observations", meta.get("valid_pe_quarters", x.get("backtest_valid_pe_quarters"))),
        ("Backtest start", meta.get("first_quarter")),
        ("Backtest end", meta.get("last_quarter")),
        ("Model", "70% own-index P/E valuation percentile + 30% own-index YoY implied EPS growth"),
        ("P/E source", "NSE Indices daily archive"),
    ]
    for j,(lab,val) in enumerate(notes, start=note_row+1):
        ws.cell(j,1,lab).font = gray_font
        ws.cell(j,2,_safe_excel_value(val))
        if lab == "P/E source":
            ws.cell(j,2).hyperlink = NSE_SOURCE_URL
            ws.cell(j,2).style = "Hyperlink"
    _autosize(ws)

    # Monthly P/E history
    wm = wb.create_sheet("Monthly PE History")
    _style_workbook_sheet(wm, "A2")
    mh = ["Date","Index Close","P/E","Implied EPS","Earnings Yield"]
    for c,h in enumerate(mh,1): wm.cell(1,c,h)
    _apply_table_header(wm,1,1,len(mh))
    _set_imported_header_comment(wm["B1"], NSE_SOURCE_URL)
    _set_imported_header_comment(wm["C1"], NSE_SOURCE_URL)
    md = monthly_detail.copy().sort_values("Date") if not monthly_detail.empty else monthly_detail
    for i,(_,row) in enumerate(md.iterrows(), start=2):
        dt = pd.to_datetime(row.get("Date"), errors="coerce")
        wm.cell(i,1,_safe_excel_value(dt)).number_format="dd-mmm-yyyy"
        wm.cell(i,2,_safe_excel_value(row.get("Close"))).number_format='#,##0.00;[Red](#,##0.00);-'
        wm.cell(i,3,_safe_excel_value(row.get("PE"))).number_format='0.00x;[Red](0.00x);-'
        wm.cell(i,2).font = green_font; wm.cell(i,3).font = green_font
        wm.cell(i,4,f'=IFERROR(B{i}/C{i},"")').font = black_font
        wm.cell(i,4).number_format='#,##0.00;[Red](#,##0.00);-'
        wm.cell(i,5,f'=IFERROR(1/C{i},"")').font = black_font
        wm.cell(i,5).number_format='0.00%'
    _autosize(wm)

    # Quarter-end backtest detail
    wq = wb.create_sheet("Quarterly Backtest")
    _style_workbook_sheet(wq, "A2")
    qh = ["Date","Index Close","P/E","Regime","Implied EPS","YoY EPS Growth","P/E Percentile","Valuation Quintile","Fwd 1Y","Fwd 3Y CAGR","Fwd 5Y CAGR","Fwd 10Y CAGR"]
    for c,h in enumerate(qh,1): wq.cell(1,c,h)
    _apply_table_header(wq,1,1,len(qh))
    _set_imported_header_comment(wq["B1"], NSE_SOURCE_URL)
    _set_imported_header_comment(wq["C1"], NSE_SOURCE_URL)
    qd = quarterly_detail.copy().sort_values("Date") if not quarterly_detail.empty else quarterly_detail
    nrows = len(qd)
    first_excel = 2
    last_excel = first_excel + nrows - 1
    for idx,(_,row) in enumerate(qd.iterrows(), start=2):
        dt = pd.to_datetime(row.get("Date"), errors="coerce")
        close_val = row.get("Close", row.get("Price"))
        pe_val = row.get("PE")
        regime_val = row.get("Regime")
        if regime_val is None or pd.isna(regime_val):
            regime_val = "Consolidated" if pd.notna(dt) and dt >= REGIME_SEAM else "Standalone"
        wq.cell(idx,1,_safe_excel_value(dt)).number_format="dd-mmm-yyyy"
        wq.cell(idx,2,_safe_excel_value(close_val)).number_format='#,##0.00;[Red](#,##0.00);-'
        wq.cell(idx,3,_safe_excel_value(pe_val)).number_format='0.00x;[Red](0.00x);-'
        wq.cell(idx,4,regime_val)
        wq.cell(idx,2).font = green_font; wq.cell(idx,3).font = green_font; wq.cell(idx,4).font = gray_font
        wq.cell(idx,5,f'=IFERROR(B{idx}/C{idx},"")').number_format='#,##0.00;[Red](#,##0.00);-'
        # YoY EPS: same quarter one year earlier, only if the PE regime is unchanged.
        if idx >= 6:
            wq.cell(idx,6,f'=IF(AND(D{idx}=D{idx-4},E{idx-4}<>""),E{idx}/E{idx-4}-1,"")')
        else:
            wq.cell(idx,6,None)
        wq.cell(idx,6).number_format='0.0%'
        if nrows:
            wq.cell(idx,7,f'=IF(C{idx}="","",(COUNTIFS($D$2:$D${last_excel},D{idx},$C$2:$C${last_excel},">0",$C$2:$C${last_excel},"<"&C{idx})+(COUNTIFS($D$2:$D${last_excel},D{idx},$C$2:$C${last_excel},C{idx})+1)/2)/COUNTIFS($D$2:$D${last_excel},D{idx},$C$2:$C${last_excel},">0"))')
        wq.cell(idx,7).number_format='0.0%'
        wq.cell(idx,8,f'=IF(G{idx}="","",IF(G{idx}<=20%,"Q1 Cheapest",IF(G{idx}<=40%,"Q2",IF(G{idx}<=60%,"Q3",IF(G{idx}<=80%,"Q4","Q5 Most Expensive")))))')
        offsets = [(9,4,1),(10,12,3),(11,20,5),(12,40,10)]
        for col, offset, years in offsets:
            target = idx + offset
            if target <= last_excel:
                if years == 1:
                    formula = f'=IFERROR(B{target}/B{idx}-1,"")'
                else:
                    formula = f'=IFERROR((B{target}/B{idx})^(1/{years})-1,"")'
                wq.cell(idx,col,formula)
            else:
                wq.cell(idx,col,None)
            wq.cell(idx,col).number_format='0.0%'
    _autosize(wq)

    # Daily score/signal history. Percentile and EPS-growth columns are pipeline inputs;
    # scores and signal are reconstructed as formulas so the workbook remains auditable.
    wh = wb.create_sheet("Score History")
    _style_workbook_sheet(wh, "A2")
    hh = ["Date","Index Close","P/E","P/E Percentile","YoY EPS Growth","Valuation Score","Growth Score","Composite Score","Signal"]
    for c,h in enumerate(hh,1): wh.cell(1,c,h)
    _apply_table_header(wh,1,1,len(hh))
    hd = history_detail.copy().sort_values("date") if not history_detail.empty else history_detail
    for i,(_,row) in enumerate(hd.iterrows(), start=2):
        dt = pd.to_datetime(row.get("date"), errors="coerce")
        wh.cell(i,1,_safe_excel_value(dt)).number_format="dd-mmm-yyyy"
        wh.cell(i,2,_safe_excel_value(row.get("close"))).number_format='#,##0.00;[Red](#,##0.00);-'
        wh.cell(i,3,_safe_excel_value(row.get("pe"))).number_format='0.00x;[Red](0.00x);-'
        wh.cell(i,4,_safe_excel_value(row.get("pe_percentile"))).number_format='0.0%'
        wh.cell(i,5,_safe_excel_value(row.get("yoy_eps_growth"))).number_format='0.0%'
        for c in (2,3,4,5): wh.cell(i,c).font = green_font
        wh.cell(i,6,f'=IF(D{i}="","",100*(1-D{i}))').number_format='0.0'
        wh.cell(i,7,f'=IF(E{i}="","",IF(E{i}<=0,20,IF(E{i}<5%,20+400*E{i},IF(E{i}<10%,40+400*(E{i}-5%),IF(E{i}<15%,60+400*(E{i}-10%),IF(E{i}<25%,80+200*(E{i}-15%),100))))))').number_format='0.0'
        wh.cell(i,8,f'=IF(OR(F{i}="",G{i}=""),"",70%*F{i}+30%*G{i})').number_format='0.0'
        wh.cell(i,9,f'=IF(H{i}="","",IF(H{i}>=70,"BUY",IF(H{i}>=40,"HOLD","SELL")))')
    _autosize(wh)

    # Source and methodology notes.
    ws2 = wb.create_sheet("Sources & Methodology")
    _style_workbook_sheet(ws2)
    _apply_section_header(ws2,1,"Sources",4)
    src = [
        ("NSE Indices daily archive", NSE_SOURCE_URL),
        ("NSE index archive files", NSE_ARCHIVE_BASE),
    ]
    for i,(name,url) in enumerate(src,start=2):
        ws2.cell(i,1,name)
        ws2.cell(i,2,url)
        ws2.cell(i,2).hyperlink=url
        ws2.cell(i,2).style="Hyperlink"
    _apply_section_header(ws2,5,"Methodology",4)
    meth = [
        "All P/E, implied EPS and return calculations are specific to the selected index.",
        "Valuation score = 100 × (1 − own-index P/E percentile).",
        "Growth score is based on YoY change in own-index implied EPS.",
        "Composite score = 70% valuation score + 30% growth score.",
        "BUY ≥ 70; HOLD 40 to <70; SELL <40.",
        "NIFTY 50 uses the canonical long-history quarter-end backtest where available; other indices use the multi-index archive history.",
    ]
    for i,t in enumerate(meth,start=6): ws2.cell(i,1,t)
    ws2.column_dimensions['A'].width=92
    ws2.column_dimensions['B'].width=55

    # Consistent number/input visual cues and print setup.
    for sh in wb.worksheets:
        sh.sheet_properties.pageSetUpPr.fitToPage = True
        sh.page_setup.fitToWidth = 1
        sh.page_setup.fitToHeight = 0
        sh.page_margins.left = 0.3
        sh.page_margins.right = 0.3
        sh.page_margins.top = 0.5
        sh.page_margins.bottom = 0.5

    out_path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(out_path)


def generate_index_workbooks(
    web_dir: Path,
    latest_map: Dict[str, Dict],
    backtests: Dict[str, Dict],
    quarter_details: Dict[str, pd.DataFrame],
    monthly: pd.DataFrame,
    history_path: Path,
):
    """Generate one downloadable XLSX per eligible dashboard index."""
    out_dir = web_dir / "downloads" / "indices"
    out_dir.mkdir(parents=True, exist_ok=True)
    for old in out_dir.glob("*.xlsx"):
        old.unlink()

    if history_path.exists():
        hist = pd.read_csv(history_path)
    else:
        hist = pd.DataFrame()

    for slug, x in latest_map.items():
        key = x.get("index_key")
        md = monthly[monthly["IndexKey"] == key].copy() if "IndexKey" in monthly.columns else pd.DataFrame()
        qd = quarter_details.get(slug, pd.DataFrame())
        hd = hist[hist["slug"] == slug].copy() if (not hist.empty and "slug" in hist.columns) else pd.DataFrame()
        out = out_dir / f"{slug}.xlsx"
        _write_index_workbook(out, x, backtests.get(slug, {}), md, qd, hd)
        print(f"Excel: wrote {out} ({x.get('index_name')})")



def _env_bool(name: str, default: bool = False) -> bool:
    value = os.getenv(name)
    if value is None:
        return default
    return str(value).strip().lower() in {"1", "true", "yes", "y", "on"}


def _safe_float(value, default: float) -> float:
    try:
        return float(value)
    except (TypeError, ValueError):
        return float(default)


def _alert_num(value, digits=2, suffix=""):
    if value is None:
        return "—"
    try:
        return f"{float(value):,.{digits}f}{suffix}"
    except (TypeError, ValueError):
        return "—"


def _alert_pct(value, digits=1, signed=False):
    if value is None:
        return "—"
    try:
        x = float(value) * 100
    except (TypeError, ValueError):
        return "—"
    if signed:
        return f"{x:+.{digits}f}%"
    return f"{x:.{digits}f}%"


def _index_alert_snapshot(slug: str, x: Dict) -> Dict:
    bounds = x.get("signal_boundaries") or {}

    def slim(key: str):
        b = bounds.get(key) or {}
        return {
            "pe": b.get("pe"),
            # Kept as index_level in alert state even though the model's internal
            # boundary field is named nifty_level for backwards compatibility.
            "index_level": b.get("nifty_level"),
            "move_from_current": b.get("move_from_current"),
        }

    return {
        "slug": slug,
        "index_name": x.get("index_name", slug),
        "as_of": x.get("as_of"),
        "signal": x.get("signal"),
        "composite_score": x.get("composite_score"),
        "index_level": x.get("nifty_close"),
        "pe": x.get("pe"),
        "buy_hold": slim("buy_hold"),
        "hold_sell": slim("hold_sell"),
    }


def _load_multi_alert_state(path: Path) -> Dict:
    try:
        if path.exists():
            data = json.loads(path.read_text(encoding="utf-8"))
            if isinstance(data, dict) and isinstance(data.get("indices"), dict):
                return data
    except Exception as e:
        print(f"Warning: could not read multi-index alert state {path}: {e}")
    return {"indices": {}}


def _write_multi_alert_state(path: Path, snapshots: Dict[str, Dict]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    payload = {
        "updated_at": datetime.utcnow().isoformat(timespec="seconds") + "Z",
        "indices": snapshots,
    }
    path.write_text(json.dumps(payload, indent=2, allow_nan=False), encoding="utf-8")


def _compare_index_alert(previous: Dict, current: Dict, pe_delta_trigger: float, index_pct_trigger: float):
    changes = []
    boundary_triggered = False
    for key, label in (("buy_hold", "BUY → HOLD"), ("hold_sell", "HOLD → SELL")):
        old = (previous or {}).get(key) or {}
        new = current.get(key) or {}
        old_pe, new_pe = old.get("pe"), new.get("pe")
        old_level, new_level = old.get("index_level"), new.get("index_level")
        pe_delta = None
        level_pct = None
        if old_pe not in (None, 0) and new_pe is not None:
            pe_delta = float(new_pe) - float(old_pe)
        if old_level not in (None, 0) and new_level is not None:
            level_pct = float(new_level) / float(old_level) - 1.0
        this_trigger = (
            (pe_delta is not None and abs(pe_delta) >= pe_delta_trigger)
            or (level_pct is not None and abs(level_pct) >= index_pct_trigger)
        )
        boundary_triggered = boundary_triggered or this_trigger
        changes.append({
            "key": key,
            "label": label,
            "old_pe": old_pe,
            "new_pe": new_pe,
            "pe_delta": pe_delta,
            "old_level": old_level,
            "new_level": new_level,
            "level_pct": level_pct,
            "triggered": this_trigger,
        })

    signal_changed = bool(
        previous
        and previous.get("signal")
        and current.get("signal")
        and previous.get("signal") != current.get("signal")
    )
    triggered = boundary_triggered or signal_changed
    return triggered, signal_changed, changes


def _trigger_description(alert: Dict) -> str:
    parts = []
    if alert.get("signal_changed"):
        prev = alert.get("previous", {}).get("signal")
        curr = alert.get("current", {}).get("signal")
        parts.append(f"Signal {prev} → {curr}")
    boundary_labels = [c["label"] for c in alert.get("changes", []) if c.get("triggered")]
    if boundary_labels:
        if len(boundary_labels) == 2:
            parts.append("Both thresholds changed")
        else:
            parts.append(f"{boundary_labels[0]} threshold changed")
    if alert.get("test") and not parts:
        parts.append("Test alert")
    return "; ".join(parts) if parts else "Threshold update"


def _dashboard_link(base_url: str, slug: str) -> str:
    if not base_url:
        return ""
    sep = "&" if "?" in base_url else "?"
    return f"{base_url}{sep}{urlencode({'index': slug})}"


def _build_multi_index_email(alerts: List[Dict], base_url: str, force_test: bool = False) -> EmailMessage:
    if not alerts:
        raise ValueError("alerts cannot be empty")

    msg = EmailMessage()
    if force_test and len(alerts) == 1:
        name = alerts[0]["current"]["index_name"]
        subject = f"Index Valuation Tracker – {name} – Test Alert"
    elif len(alerts) == 1:
        a = alerts[0]
        name = a["current"]["index_name"]
        if a.get("signal_changed"):
            subject_tail = f"Signal changed to {a['current'].get('signal')}"
        else:
            subject_tail = "Threshold Alert"
        subject = f"Index Valuation Tracker – {name} – {subject_tail}"
    else:
        subject = f"Index Valuation Tracker – {len(alerts)} Index Alerts"
    msg["Subject"] = subject

    summary_rows = []
    for a in alerts:
        c = a["current"]
        link = _dashboard_link(base_url, c["slug"])
        name = html.escape(str(c["index_name"]))
        if link:
            name = f'<a href="{html.escape(link)}" style="color:#2457a7;text-decoration:none">{name}</a>'
        summary_rows.append(
            "<tr>"
            f"<td>{name}</td>"
            f"<td><b>{html.escape(str(c.get('signal') or '—'))}</b></td>"
            f"<td>{_alert_num(c.get('composite_score'),1)}</td>"
            f"<td>{_alert_num(c.get('pe'),2,'x')}</td>"
            f"<td>{html.escape(_trigger_description(a))}</td>"
            "</tr>"
        )

    detail_sections = []
    for a in alerts:
        c = a["current"]
        link = _dashboard_link(base_url, c["slug"])
        rows = []
        for ch in a.get("changes", []):
            row_style = "background:#fff7e6;" if ch.get("triggered") else ""
            rows.append(
                f'<tr style="{row_style}">'
                f"<td>{html.escape(ch['label'])}</td>"
                f"<td>{_alert_num(ch.get('old_pe'),2,'x')}</td>"
                f"<td><b>{_alert_num(ch.get('new_pe'),2,'x')}</b></td>"
                f"<td>{_alert_num(ch.get('old_level'),0)}</td>"
                f"<td><b>{_alert_num(ch.get('new_level'),0)}</b></td>"
                f"<td>{_alert_pct(ch.get('level_pct'),1,signed=True)}</td>"
                "</tr>"
            )
        link_html = f'<p><a href="{html.escape(link)}">Open {html.escape(c["index_name"])} dashboard</a></p>' if link else ""
        signal_note = ""
        if a.get("signal_changed"):
            signal_note = (
                f'<p style="padding:9px 11px;background:#eef5ff;border-left:3px solid #2457a7">'
                f'<b>Signal changed:</b> {html.escape(str(a.get("previous",{}).get("signal") or "—"))} '
                f'→ <b>{html.escape(str(c.get("signal") or "—"))}</b></p>'
            )
        detail_sections.append(f"""
          <div style="margin-top:28px">
            <h2 style="font-size:18px;margin:0 0 10px">{html.escape(c['index_name'])} Valuation Signal Alert</h2>
            <p><b>As of:</b> {html.escape(str(c.get('as_of') or '—'))} &nbsp; | &nbsp;
               <b>Index Level:</b> {_alert_num(c.get('index_level'),2)} &nbsp; | &nbsp;
               <b>P/E:</b> {_alert_num(c.get('pe'),2,'x')}</p>
            <p><b>Composite score:</b> {_alert_num(c.get('composite_score'),1)} &nbsp; | &nbsp;
               <b>Signal:</b> {html.escape(str(c.get('signal') or '—'))}</p>
            {signal_note}
            <table cellpadding="7" cellspacing="0" border="1" style="border-collapse:collapse;border-color:#d8dee8;width:100%;max-width:760px">
              <tr style="background:#f5f7fa"><th>Boundary</th><th>Previous P/E</th><th>Current P/E</th><th>Previous Index Level</th><th>Current Index Level</th><th>Level change</th></tr>
              {''.join(rows)}
            </table>
            <p style="color:#5f6f82">Threshold index levels assume implied EPS and the earnings-growth score remain unchanged.</p>
            {link_html}
          </div>
        """)

    summary_heading = "Test alert" if force_test else ("Alert summary" if len(alerts) > 1 else "Alert")
    html_body = f"""
    <html><body style="font-family:Arial,sans-serif;color:#142033;line-height:1.45">
      <h1 style="font-size:22px;margin-bottom:6px">Index Valuation Tracker</h1>
      <p style="color:#5f6f82;margin-top:0">{summary_heading}</p>
      <table cellpadding="7" cellspacing="0" border="1" style="border-collapse:collapse;border-color:#d8dee8;width:100%;max-width:760px">
        <tr style="background:#f5f7fa"><th>Index</th><th>Signal</th><th>Score</th><th>P/E</th><th>Trigger</th></tr>
        {''.join(summary_rows)}
      </table>
      {''.join(detail_sections)}
    </body></html>
    """

    plain = ["Index Valuation Tracker", ""]
    for a in alerts:
        c = a["current"]
        plain += [
            f"{c['index_name']} — {_trigger_description(a)}",
            f"As of: {c.get('as_of')}",
            f"Index Level: {_alert_num(c.get('index_level'),2)} | P/E: {_alert_num(c.get('pe'),2,'x')}",
            f"Composite score: {_alert_num(c.get('composite_score'),1)} | Signal: {c.get('signal')}",
        ]
        for ch in a.get("changes", []):
            plain.append(
                f"{ch['label']}: P/E {_alert_num(ch.get('old_pe'),2,'x')} -> {_alert_num(ch.get('new_pe'),2,'x')}; "
                f"Index {_alert_num(ch.get('old_level'),0)} -> {_alert_num(ch.get('new_level'),0)} "
                f"({_alert_pct(ch.get('level_pct'),1,signed=True)})"
            )
        link = _dashboard_link(base_url, c["slug"])
        if link:
            plain.append(link)
        plain.append("")

    msg.set_content("\n".join(plain))
    msg.add_alternative(html_body, subtype="html")
    return msg


def _send_multi_index_email(alerts: List[Dict], base_url: str, force_test: bool = False) -> bool:
    to_addr = os.getenv("ALERT_EMAIL_TO", "").strip()
    username = os.getenv("SMTP_USERNAME", "").strip()
    password = os.getenv("SMTP_PASSWORD", "").strip()
    if not (to_addr and username and password):
        print("Multi-index email not sent: configure ALERT_EMAIL_TO, SMTP_USERNAME and SMTP_PASSWORD.")
        return False

    host = os.getenv("SMTP_HOST", "smtp.gmail.com").strip() or "smtp.gmail.com"
    port = int(os.getenv("SMTP_PORT", "465"))
    from_addr = os.getenv("SMTP_FROM", username).strip() or username
    msg = _build_multi_index_email(alerts, base_url, force_test=force_test)
    msg["From"] = from_addr
    msg["To"] = to_addr

    try:
        with smtplib.SMTP_SSL(host, port, timeout=30) as smtp:
            smtp.login(username, password)
            smtp.send_message(msg)
        print(f"Multi-index alert email sent to {to_addr}: {msg['Subject']}")
        return True
    except Exception as e:
        print(f"Warning: multi-index alert email failed: {e}")
        return False


def process_multi_index_alerts(web_dir: Path, latest_map: Dict[str, Dict]) -> None:
    """Send one index-specific or consolidated email for all meaningful alert events."""
    if not latest_map:
        return

    state_path = web_dir / "data" / "multi_index_alert_state.json"
    prior_state = _load_multi_alert_state(state_path).get("indices", {})
    current_state = {slug: _index_alert_snapshot(slug, x) for slug, x in latest_map.items()}

    force = _env_bool("FORCE_THRESHOLD_ALERT", False)
    pe_delta_trigger = _safe_float(os.getenv("ALERT_PE_DELTA"), 0.10)
    index_pct_trigger = _safe_float(
        os.getenv("ALERT_INDEX_DELTA_PCT", os.getenv("ALERT_NIFTY_DELTA_PCT")), 0.005
    )
    base_url = os.getenv("DASHBOARD_URL", "").strip()

    # First run: establish baselines. A manual test still sends a deterministic
    # NIFTY 50 example so the user can validate subject, body and dashboard link.
    if not prior_state:
        if force:
            test_slug = next((s for s,x in latest_map.items() if x.get("index_name") == "NIFTY 50"), next(iter(latest_map)))
            c = current_state[test_slug]
            dummy_changes = []
            for key, label in (("buy_hold", "BUY → HOLD"), ("hold_sell", "HOLD → SELL")):
                b = c.get(key) or {}
                dummy_changes.append({
                    "key": key, "label": label,
                    "old_pe": b.get("pe"), "new_pe": b.get("pe"), "pe_delta": 0.0,
                    "old_level": b.get("index_level"), "new_level": b.get("index_level"), "level_pct": 0.0,
                    "triggered": False,
                })
            _send_multi_index_email([{
                "previous": c, "current": c, "signal_changed": False,
                "changes": dummy_changes, "test": True,
            }], base_url, force_test=True)
        _write_multi_alert_state(state_path, current_state)
        print("Multi-index alert baseline initialized.")
        return

    alerts = []
    for slug, c in current_state.items():
        p = prior_state.get(slug)
        # Newly eligible indices get a baseline without generating alert spam.
        if not p:
            continue
        triggered, signal_changed, changes = _compare_index_alert(p, c, pe_delta_trigger, index_pct_trigger)
        if triggered:
            alerts.append({
                "previous": p,
                "current": c,
                "signal_changed": signal_changed,
                "changes": changes,
                "test": False,
            })

    if force and not alerts:
        test_slug = next((s for s,x in latest_map.items() if x.get("index_name") == "NIFTY 50"), next(iter(latest_map)))
        c = current_state[test_slug]
        p = prior_state.get(test_slug, c)
        _, signal_changed, changes = _compare_index_alert(p, c, pe_delta_trigger, index_pct_trigger)
        # Make test output easy to read even when no values moved.
        if not changes:
            changes = []
        alerts = [{
            "previous": p,
            "current": c,
            "signal_changed": signal_changed,
            "changes": changes,
            "test": True,
        }]

    if alerts:
        test_mode = bool(force and all(a.get("test") for a in alerts))
        sent = _send_multi_index_email(alerts, base_url, force_test=test_mode)
        email_configured = bool(
            os.getenv("ALERT_EMAIL_TO") and os.getenv("SMTP_USERNAME") and os.getenv("SMTP_PASSWORD")
        )
        if sent or not email_configured or force:
            _write_multi_alert_state(state_path, current_state)
        else:
            # Preserve the old baseline so the same real alert is retried next run.
            print("Multi-index alert state retained because email delivery failed.")
    else:
        _write_multi_alert_state(state_path, current_state)
        print(
            f"No multi-index alert: changes are below {pe_delta_trigger:.2f}x P/E / "
            f"{index_pct_trigger:.1%} index-level tolerances and no signals changed."
        )

def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--web-dir", default="docs")
    ap.add_argument("--quarterly-csv", default="multi_index_quarterly.csv")
    ap.add_argument("--monthly-csv", default="multi_index_monthly.csv")
    ap.add_argument("--target-date", default=None, help="YYYY-MM-DD; defaults to today")
    ap.add_argument(
        "--nifty50-quarterly-csv",
        default="nifty_quarterly_backtest.csv",
        help="Canonical long-history NIFTY 50 quarter-end backtest file (1999 onward)",
    )
    args = ap.parse_args()

    web_dir = Path(args.web_dir)
    data_dir = web_dir / "data"
    qpath = Path(args.quarterly_csv)
    mpath = Path(args.monthly_csv)
    target = date.fromisoformat(args.target_date) if args.target_date else date.today()
    nifty50_legacy_path = Path(args.nifty50_quarterly_csv)
    nifty50_legacy_summary = build_legacy_nifty50_backtest(nifty50_legacy_path)
    nifty50_legacy_detail = load_legacy_nifty50_detail(nifty50_legacy_path)

    s = session()
    latest_date, latest_snap, latest_url = fetch_snapshot_on_or_before(s, target, max_lookback=12)
    # Approximately one year before the same observation date; last NSE trading day on/before.
    prior_target = date(latest_date.year - 1, latest_date.month, min(latest_date.day, monthrange(latest_date.year - 1, latest_date.month)[1]))
    prior_date, prior_snap, _ = fetch_snapshot_on_or_before(s, prior_target, max_lookback=12)

    q_end = latest_completed_quarter_end(latest_date)
    m_end = latest_completed_month_end(latest_date)
    q_periods = quarter_ends(BOOTSTRAP_QUARTER_START, q_end)
    m_periods = month_ends(BOOTSTRAP_MONTH_START, m_end)

    quarterly = append_period_snapshots(s, qpath, q_periods, "Quarterly")
    monthly = append_period_snapshots(s, mpath, m_periods, "Monthly")

    # Current universe is the equity-like rows present in latest NSE snapshot.
    # Preserve all such indices even if current PE is unavailable; website will flag them.
    current = latest_snap.copy().sort_values(["Group","IndexName"])

    latest_map: Dict[str, Dict] = {}
    backtests: Dict[str, Dict] = {}
    quarter_details: Dict[str, pd.DataFrame] = {}
    used_slugs = set()
    for _, row in current.iterrows():
        key = row["IndexKey"]
        qbt, summary = build_quarterly_backtest(quarterly, key)
        # NIFTY 50 is special: preserve the original 1999-onward canonical
        # backtest instead of the much shorter generic multi-index archive.
        # Current/live valuation metrics still come from the multi-index engine;
        # only the historical forward-return table and its sample metadata are
        # replaced by the deeper NIFTY 50 dataset.
        if key == "NIFTY50" and nifty50_legacy_summary is not None:
            summary = nifty50_legacy_summary
            print(
                "NIFTY 50: using canonical long-history backtest "
                f"{summary['meta'].get('first_quarter')} to {summary['meta'].get('last_quarter')} "
                f"({summary['meta'].get('valid_pe_quarters')} PE-bearing quarters)"
            )
        base_slug = slugify(row["IndexName"])
        slug = base_slug
        n = 2
        while slug in used_slugs:
            slug = f"{base_slug}-{n}"; n += 1
        used_slugs.add(slug)

        x = build_current_for_index(key, row, latest_date, prior_snap, monthly, qbt)
        if key == "NIFTY50" and nifty50_legacy_summary is not None:
            x["backtest_valid_pe_quarters"] = int(nifty50_legacy_summary["meta"].get("valid_pe_quarters", 0) or 0)
        x["slug"] = slug
        latest_map[slug] = x
        backtests[slug] = summary
        if key == "NIFTY50" and nifty50_legacy_detail is not None:
            quarter_details[slug] = nifty50_legacy_detail.copy()
        else:
            quarter_details[slug] = qbt.copy()

    # Dashboard eligibility: show only indices for which the model is actually useful.
    # This removes "limited history" and live P/E-unavailable entries from the
    # selector/heatmap rather than displaying a weak or incomplete signal.
    #
    # Requirements:
    #   * current P/E is available
    #   * live composite score can be calculated
    #   * at least MIN_LIVE_MONTHS of comparable post-2021 P/E history
    #   * at least MIN_BACKTEST_QUARTERS PE-bearing quarter-end observations
    #   * at least MIN_MATURE_3Y_OBS matured 3-year forward-return observations
    #
    # This deliberately favours a smaller, more statistically useful dashboard
    # universe over showing every NSE index with only a short back-cast history.
    #
    # The full raw monthly/quarterly archive is still retained, so a currently
    # excluded index can automatically enter the dashboard once it matures.
    keep = set()
    excluded = []
    for slug, x in latest_map.items():
        bt_meta = backtests.get(slug, {}).get("meta", {})
        q_obs = int(bt_meta.get("valid_pe_quarters", 0) or 0)
        matured_3y_obs = sum(int(r.get("n_3y", 0) or 0) for r in backtests.get(slug, {}).get("rows", []))
        live_months = int(x.get("live_pe_history_months", 0) or 0)
        useful = (
            x.get("pe") is not None
            and x.get("composite_score") is not None
            and live_months >= MIN_LIVE_MONTHS
            and q_obs >= MIN_BACKTEST_QUARTERS
            and matured_3y_obs >= MIN_MATURE_3Y_OBS
        )
        if useful:
            keep.add(slug)
        else:
            excluded.append((
                x.get("index_name", slug), live_months, q_obs, matured_3y_obs,
                x.get("pe") is not None
            ))

    latest_map = {k:v for k,v in latest_map.items() if k in keep}
    backtests = {k:v for k,v in backtests.items() if k in keep}
    quarter_details = {k:v for k,v in quarter_details.items() if k in keep}

    if excluded:
        print(f"Excluded {len(excluded)} limited/unusable indices from dashboard.")
        for name, live_months, q_obs, matured_3y_obs, has_pe in excluded:
            print(
                f"  - {name}: live_months={live_months}, quarter_pe_obs={q_obs}, "
                f"matured_3y_obs={matured_3y_obs}, current_pe={'yes' if has_pe else 'no'}"
            )

    catalog = make_catalog(latest_map, backtests)
    data_dir.mkdir(parents=True, exist_ok=True)
    (data_dir / "multi_index_catalog.json").write_text(json.dumps(catalog, indent=2, allow_nan=False), encoding="utf-8")
    (data_dir / "multi_index_latest.json").write_text(json.dumps({
        "as_of": latest_date.isoformat(),
        "prior_as_of": prior_date.isoformat(),
        "source_url": latest_url,
        "indices": latest_map,
    }, indent=2, allow_nan=False), encoding="utf-8")
    (data_dir / "multi_index_backtests.json").write_text(json.dumps({
        "generated_at": datetime.utcnow().isoformat(timespec="seconds") + "Z",
        "indices": backtests,
    }, indent=2, allow_nan=False), encoding="utf-8")
    history_path = data_dir / "multi_index_history.csv"
    append_daily_history(history_path, latest_map)
    generate_index_workbooks(web_dir, latest_map, backtests, quarter_details, monthly, history_path)
    process_multi_index_alerts(web_dir, latest_map)

    ok = sum(1 for x in latest_map.values() if x.get("signal"))
    print(f"Multi-index tracker updated: {len(latest_map)} eligible indices; {ok} with live signal; as of {latest_date}")


if __name__ == "__main__":
    main()
