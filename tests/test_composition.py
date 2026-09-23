import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch

from openpyxl import load_workbook, Workbook

from composition_tracker import (
    discover_index_pages, fetch_constituent_file, parse_constituent_csv,
    update_workbook_composition,
)


class CompositionSourcesTest(unittest.TestCase):
    item = {"slug": "nifty-100", "name": "NIFTY 100"}
    url = "https://www.niftyindices.com/IndexConstituent/ind_nifty100list.csv"

    def test_official_list_has_members_and_no_invented_weights(self):
        csv = ("Company Name,Industry,Symbol,Series,ISIN Code\n"
               "Acme Bank Ltd.,Banking,ACME,EQ,INE001\n"
               "Beta Motors Ltd.,Automobile,BETA,EQ,INE002\n"
               "Gamma Bank Ltd.,Banking,GAMMA,EQ,INE003\n").encode()
        data = parse_constituent_csv(csv, self.item, self.url, "2026-09-23")
        self.assertEqual(data["stock_count"], 3)
        self.assertIsNone(data["weight_coverage"])
        self.assertIsNone(data["top10_weight"])
        self.assertTrue(all(row["weight"] is None for row in data["holdings"]))
        self.assertEqual(data["sectors"][0]["name"], "Banking")
        self.assertEqual(data["sectors"][0]["count"], 2)
        self.assertIsNone(data["as_of"])
        self.assertEqual(data["retrieved_on"], "2026-09-23")

        with tempfile.TemporaryDirectory() as tmp:
            book = Path(tmp) / "nifty-100.xlsx"
            Workbook().save(book)
            update_workbook_composition(book, data)
            ws = load_workbook(book)["Composition"]
            self.assertEqual(ws["E2"].value, "Not published")
            self.assertIsNone(ws["E6"].value)

    def test_discovery_follows_only_the_official_index_page_and_csv(self):
        category = ("<a href='/indices/equity/broad-based-indices/nifty-100'>Nifty 100</a>"
                    "<a href='https://elsewhere.example/indices/equity/broad-based-indices/nifty-100'>Nifty 100</a>")
        detail = ("<a href='https://elsewhere.example/IndexConstituent/bad.csv'>Index Constituent</a>"
                  "<a href='/IndexConstituent/ind_nifty100list.csv'>Index Constituent</a>")
        csv = ("Company Name,Industry,Symbol\nAcme Bank Ltd.,Banking,ACME\n"
               "Beta Motors Ltd.,Automobile,BETA\nGamma Bank Ltd.,Banking,GAMMA\n")
        def load(url):
            if url.endswith("nifty-100"):
                return detail.encode()
            if url.endswith(".csv"):
                return csv.encode()
            return category.encode()
        with patch("composition_tracker.official_get", side_effect=load):
            pages = discover_index_pages({"items": [self.item]})
            self.assertEqual(pages[self.item["slug"]],
                             "https://www.niftyindices.com/indices/equity/broad-based-indices/nifty-100")
            data = fetch_constituent_file(self.item, pages[self.item["slug"]], "2026-09-23")
        self.assertEqual(data["source_url"], self.url)
        self.assertEqual(len(data["holdings"]), 3)

    def test_weights_are_used_only_when_the_whole_file_is_consistent(self):
        csv = ("Company Name,Industry,Symbol,Weight (%)\n"
               "Acme Bank,Banking,ACME,50\n"
               "Beta Motors,Automobile,BETA,30\n"
               "Gamma Bank,Banking,GAMMA,20\n").encode()
        data = parse_constituent_csv(csv, self.item, self.url, "2026-09-23")
        self.assertEqual(data["weight_coverage"], 100)
        self.assertEqual(data["top10_weight"], 100)
        partial = csv.replace(b"GAMMA,20", b"GAMMA,")
        data = parse_constituent_csv(partial, self.item, self.url, "2026-09-23")
        self.assertIsNone(data["weight_coverage"])
        self.assertTrue(all(row["weight"] is None for row in data["holdings"]))


if __name__ == "__main__":
    unittest.main()
