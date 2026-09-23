#!/usr/bin/env python3
"""Refresh official NSE/Nifty index constituent composition + weights.

Data source
-----------
NSE Indices monthly report: "Indices Market Capitalisation & Weightage".
The public report is exposed as a ZIP whose filename follows the pattern
``indices_data{Mon}{YYYY}.zip``. The archive can contain PDF/CSV/XLS/XLSX
files. This script is deliberately format-tolerant and parses whichever
representation NSE publishes for that month.

Outputs
-------
- docs/data/composition/<slug>.json
- docs/data/composition/manifest.json
- docs/data/composition/status.json
- adds/replaces a "Composition" worksheet in docs/downloads/indices/*.xlsx

The script never fabricates weights. If an index cannot be parsed from the
official report, it is marked unavailable and the dashboard shows a clear
message. Existing *official* files are retained if a transient download/parsing
failure occurs; old sample/demo files are removed.
"""
from __future__ import annotations

import argparse
import calendar
import io
import json
import os
import re
import sys
import tempfile
import zipfile
from dataclasses import dataclass
from datetime import date, datetime, timezone
from pathlib import Path
from typing import Dict, Iterable, List, Optional, Tuple

import pandas as pd
import requests
from openpyxl import load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

try:
    import pdfplumber
except Exception:
    pdfplumber = None

try:
    import cloudscraper
except Exception:
    cloudscraper = None

REPORT_PAGE = "https://www.niftyindices.com/reports/monthly-reports"
REPORT_BASES = [
    "https://www.niftyindices.com/Indices_-_Market_Capitalisation_and_Weightage",
    "https://niftyindices.com/Indices_-_Market_Capitalisation_and_Weightage",
]
MONTH_ABBR = ["Jan", "Feb", "Mar", "Apr", "May", "Jun", "Jul", "Aug", "Sep", "Oct", "Nov", "Dec"]

COMPANY_KEYS = ("company", "security", "constituent", "stock", "issuer", "name of security")
SYMBOL_KEYS = ("symbol", "ticker")
SECTOR_KEYS = ("sector", "industry", "basic industry")
WEIGHT_KEYS = ("weight", "weightage")
INDEX_KEYS = ("index name", "index", "benchmark")


def norm_space(v) -> str:
    return re.sub(r"\s+", " ", str(v or "").replace("\n", " ").strip())


def compact(v) -> str:
    return re.sub(r"[^A-Z0-9]", "", norm_space(v).upper())


def clean_header(v) -> str:
    return norm_space(v).lower().replace("%", " percent ")


def clean_company(v) -> str:
    s = norm_space(v)
    s = re.sub(r"^\s*\d+[\.)-]?\s+", "", s)
    return s.strip(" -|:")


def parse_weight(v) -> Optional[float]:
    if v is None:
        return None
    if isinstance(v, (int, float)) and pd.notna(v):
        return float(v)
    s = norm_space(v).replace(",", "")
    if not s or s.lower() in {"nan", "na", "n/a", "-", "—"}:
        return None
    s = s.replace("%", "")
    m = re.search(r"[-+]?\d+(?:\.\d+)?", s)
    if not m:
        return None
    try:
        return float(m.group())
    except ValueError:
        return None


def safe_json(path: Path):
    try:
        return json.loads(path.read_text(encoding="utf-8"))
    except Exception:
        return None


def atomic_json(path: Path, data: dict):
    path.parent.mkdir(parents=True, exist_ok=True)
    tmp = path.with_suffix(path.suffix + ".tmp")
    tmp.write_text(json.dumps(data, indent=2, ensure_ascii=False), encoding="utf-8")
    tmp.replace(path)


def catalog_index_map(catalog: dict) -> Dict[str, dict]:
    return {i["slug"]: i for i in catalog.get("items", [])}


def title_matcher(catalog: dict):
    items = []
    for item in catalog.get("items", []):
        c = compact(item.get("name"))
        if c:
            items.append((len(c), c, item["slug"], item.get("name")))
    items.sort(reverse=True)

    aliases = {
        "NIFTYMIDSMALLCAP400": "nifty-midsmallcap-400",
        "NIFTY100ESG": "nifty100-esg",
        "NIFTY100ENHANCEDESG": "nifty100-enhanced-esg",
        "NIFTY100EQUALWEIGHT": "nifty100-equal-weight",
        "NIFTY100LOWVOLATILITY30": "nifty100-low-volatility-30",
        "NIFTY100QUALITY30": "nifty100-quality-30",
        "NIFTY200QUALITY30": "nifty200-quality-30",
        "NIFTY50EQUALWEIGHT": "nifty50-equal-weight",
        "NIFTY50SHARIAH": "nifty50-shariah",
        "NIFTY50VALUE20": "nifty50-value-20",
        "NIFTY500SHARIAH": "nifty500-shariah",
        "NIFTYSMEEMERGE": "nifty-sme-emerge",
    }

    def match(text: str) -> Optional[str]:
        c = compact(text)
        if not c:
            return None
        for a, slug in aliases.items():
            if a in c:
                return slug
        for _, key, slug, _ in items:
            if key in c:
                return slug
        return None

    return match


def browser_headers() -> dict:
    return {
        "User-Agent": "Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 Chrome/128 Safari/537.36",
        "Accept": "application/zip,application/octet-stream,*/*",
        "Accept-Language": "en-US,en;q=0.9",
        "Referer": REPORT_PAGE,
        "Connection": "keep-alive",
    }


def candidate_months(today: date, count: int = 8) -> Iterable[Tuple[int, int]]:
    y, m = today.year, today.month
    for _ in range(count):
        yield y, m
        m -= 1
        if m == 0:
            m = 12
            y -= 1


def is_zip_bytes(content: bytes) -> bool:
    try:
        return zipfile.is_zipfile(io.BytesIO(content))
    except Exception:
        return False


def fetch_report_zip(today: date, max_months: int = 8) -> Tuple[bytes, str, str]:
    sessions = []
    s = requests.Session()
    s.headers.update(browser_headers())
    sessions.append(("requests", s))
    if cloudscraper is not None:
        try:
            cs = cloudscraper.create_scraper(browser={"browser": "chrome", "platform": "windows", "mobile": False})
            cs.headers.update(browser_headers())
            sessions.append(("cloudscraper", cs))
        except Exception:
            pass

    errors = []
    for year, month in candidate_months(today, max_months):
        mon = MONTH_ABBR[month - 1]
        filename = f"indices_data{mon}{year}.zip"
        for base in REPORT_BASES:
            url = f"{base}/{filename}"
            for label, sess in sessions:
                try:
                    r = sess.get(url, timeout=35, allow_redirects=True)
                    if r.ok and is_zip_bytes(r.content):
                        return r.content, url, f"{mon} {year}"
                    errors.append(f"{label} {url}: HTTP {r.status_code}, {len(r.content)} bytes")
                except Exception as e:
                    errors.append(f"{label} {url}: {e}")
    raise RuntimeError("Unable to download an official monthly composition report. " + " | ".join(errors[-8:]))


def find_header_row(raw: pd.DataFrame, scan: int = 18) -> Optional[int]:
    for r in range(min(scan, len(raw))):
        vals = [clean_header(x) for x in raw.iloc[r].tolist()]
        if any(any(k in v for k in WEIGHT_KEYS) for v in vals):
            return r
    return None


def frame_with_header(raw: pd.DataFrame) -> Optional[pd.DataFrame]:
    if raw is None or raw.empty:
        return None
    raw = raw.dropna(axis=1, how="all").dropna(axis=0, how="all")
    if raw.empty:
        return None
    hr = find_header_row(raw)
    if hr is None:
        # If pandas already supplied meaningful columns, keep them.
        cols = [clean_header(x) for x in raw.columns]
        if any(any(k in c for k in WEIGHT_KEYS) for c in cols):
            df = raw.copy()
        else:
            return None
    else:
        cols = [norm_space(x) or f"col_{i}" for i, x in enumerate(raw.iloc[hr].tolist())]
        df = raw.iloc[hr + 1 :].copy()
        df.columns = cols
    df = df.dropna(axis=0, how="all").dropna(axis=1, how="all")
    return df if not df.empty else None


def pick_col(columns: List[str], keys: Tuple[str, ...]) -> Optional[str]:
    for c in columns:
        h = clean_header(c)
        if any(k in h for k in keys):
            return c
    return None


def dataframe_records(df: pd.DataFrame, default_slug: Optional[str], match_title, source_hint: str) -> Dict[str, List[dict]]:
    out: Dict[str, List[dict]] = {}
    f = frame_with_header(df)
    if f is None:
        return out
    cols = list(f.columns)
    wcol = pick_col(cols, WEIGHT_KEYS)
    ccol = pick_col(cols, COMPANY_KEYS)
    scol = pick_col(cols, SYMBOL_KEYS)
    seccol = pick_col(cols, SECTOR_KEYS)
    icol = pick_col(cols, INDEX_KEYS)
    if not wcol or not ccol:
        return out

    for _, row in f.iterrows():
        company = clean_company(row.get(ccol))
        if not company or company.lower().startswith("total"):
            continue
        weight = parse_weight(row.get(wcol))
        if weight is None or weight <= 0 or weight > 100:
            continue
        slug = default_slug
        if icol and pd.notna(row.get(icol)):
            slug = match_title(str(row.get(icol))) or slug
        if not slug:
            continue
        rec = {
            "name": company,
            "symbol": norm_space(row.get(scol)) if scol and pd.notna(row.get(scol)) else "",
            "sector": norm_space(row.get(seccol)) if seccol and pd.notna(row.get(seccol)) else "",
            "weight": round(float(weight), 6),
            "source_hint": source_hint,
        }
        out.setdefault(slug, []).append(rec)
    return out


def merge_records(target: Dict[str, List[dict]], incoming: Dict[str, List[dict]]):
    for slug, rows in incoming.items():
        target.setdefault(slug, []).extend(rows)


def parse_csv_bytes(data: bytes, filename: str, match_title) -> Dict[str, List[dict]]:
    for enc in ("utf-8-sig", "utf-8", "cp1252", "latin1"):
        try:
            text = data.decode(enc)
            raw = pd.read_csv(io.StringIO(text), header=None, dtype=object, engine="python")
            slug = match_title(filename + " " + " ".join(map(str, raw.head(8).fillna("").values.flatten())))
            return dataframe_records(raw, slug, match_title, filename)
        except Exception:
            continue
    return {}


def parse_excel_bytes(data: bytes, filename: str, match_title) -> Dict[str, List[dict]]:
    out: Dict[str, List[dict]] = {}
    for engine in (None, "openpyxl", "xlrd"):
        try:
            xls = pd.ExcelFile(io.BytesIO(data), engine=engine)
            for sheet in xls.sheet_names:
                raw = pd.read_excel(xls, sheet_name=sheet, header=None, dtype=object)
                intro = " ".join(map(str, raw.head(8).fillna("").values.flatten()))
                slug = match_title(f"{filename} {sheet} {intro}")
                merge_records(out, dataframe_records(raw, slug, match_title, f"{filename}:{sheet}"))
            if out:
                return out
        except Exception:
            continue
    return out


def pdf_date(text: str) -> Optional[str]:
    if not text:
        return None
    for pattern in (
        r"\b(January|February|March|April|May|June|July|August|September|October|November|December)\s+(\d{1,2}),\s+(20\d{2})\b",
        r"\b(\d{1,2})[-\s](Jan|Feb|Mar|Apr|May|Jun|Jul|Aug|Sep|Oct|Nov|Dec)[-\s](20\d{2})\b",
    ):
        m = re.search(pattern, text, flags=re.I)
        if m:
            try:
                if m.lastindex == 3 and m.group(1).isalpha() and len(m.group(1)) > 3:
                    dt = datetime.strptime(f"{m.group(1)} {m.group(2)} {m.group(3)}", "%B %d %Y")
                else:
                    dt = datetime.strptime(f"{m.group(1)} {m.group(2)} {m.group(3)}", "%d %b %Y")
                return dt.date().isoformat()
            except Exception:
                pass
    return None



def pdf_word_table_records(page, default_slug: Optional[str], filename: str) -> Dict[str, List[dict]]:
    """Extract rows from PDF text using header x-positions.

    This is a fallback for NSE PDFs that visually contain a table but do not
    expose ruled lines to pdfplumber's default table extractor.
    """
    out: Dict[str, List[dict]] = {}
    try:
        words = page.extract_words(use_text_flow=True, keep_blank_chars=False) or []
    except Exception:
        return out
    if not words:
        return out

    # Group words by visual line (top coordinate).
    lines: List[List[dict]] = []
    for w in sorted(words, key=lambda z: (round(float(z.get("top", 0)), 1), float(z.get("x0", 0)))):
        top = float(w.get("top", 0))
        target = None
        for line in reversed(lines[-4:]):
            if abs(float(line[0].get("top", 0)) - top) <= 2.2:
                target = line
                break
        if target is None:
            target = []
            lines.append(target)
        target.append(w)

    for i, line in enumerate(lines):
        joined = " ".join(w.get("text", "") for w in line)
        low = clean_header(joined)
        if not any(k in low for k in WEIGHT_KEYS):
            continue
        if not any(k in low for k in COMPANY_KEYS):
            continue

        # Infer column starts from header tokens.
        tokens = sorted(line, key=lambda z: float(z.get("x0", 0)))
        def token_x(keys):
            for w in tokens:
                t = clean_header(w.get("text", ""))
                if any(k in t for k in keys):
                    return float(w.get("x0", 0))
            return None

        company_x = token_x(("company", "security", "constituent", "stock", "issuer", "name"))
        sector_x = token_x(("industry", "sector"))
        symbol_x = token_x(("symbol", "ticker"))
        weight_x = token_x(("weight", "weightage"))
        if weight_x is None:
            continue
        if company_x is None:
            company_x = min(float(w.get("x0", 0)) for w in tokens)

        # Read subsequent visual lines while they look like data rows.
        misses = 0
        for row_line in lines[i + 1 :]:
            row_top = float(row_line[0].get("top", 0))
            # Stop at obvious new section headers after a couple of misses.
            row_text = " ".join(w.get("text", "") for w in row_line).strip()
            if not row_text:
                continue
            cols = {"company": [], "sector": [], "symbol": [], "weight": []}
            for w in sorted(row_line, key=lambda z: float(z.get("x0", 0))):
                x = float(w.get("x0", 0))
                txt = w.get("text", "")
                if x >= weight_x - 2:
                    cols["weight"].append(txt)
                elif symbol_x is not None and x >= symbol_x - 2:
                    cols["symbol"].append(txt)
                elif sector_x is not None and x >= sector_x - 2:
                    cols["sector"].append(txt)
                else:
                    cols["company"].append(txt)
            wt = parse_weight(" ".join(cols["weight"]))
            company = clean_company(" ".join(cols["company"]))
            if wt is None or not company or wt <= 0 or wt > 100:
                misses += 1
                if misses >= 3:
                    break
                continue
            misses = 0
            rec = {
                "name": company,
                "symbol": norm_space(" ".join(cols["symbol"])),
                "sector": norm_space(" ".join(cols["sector"])),
                "weight": round(float(wt), 6),
                "source_hint": filename,
            }
            if default_slug:
                out.setdefault(default_slug, []).append(rec)
        # One constituent table per page is normally sufficient.
        if out:
            break
    return out

def parse_pdf_bytes(data: bytes, filename: str, match_title) -> Tuple[Dict[str, List[dict]], Dict[str, str]]:
    out: Dict[str, List[dict]] = {}
    dates: Dict[str, str] = {}
    if pdfplumber is None:
        return out, dates
    try:
        with pdfplumber.open(io.BytesIO(data)) as pdf:
            for page in pdf.pages:
                text = page.extract_text() or ""
                slug = match_title(f"{filename} {text[:2500]}")
                if slug:
                    dt = pdf_date(text)
                    if dt:
                        dates[slug] = dt
                tables = []
                try:
                    tables = page.extract_tables() or []
                except Exception:
                    tables = []
                page_before = sum(len(v) for v in out.values())
                for tbl in tables:
                    if not tbl or len(tbl) < 2:
                        continue
                    raw = pd.DataFrame(tbl)
                    merge_records(out, dataframe_records(raw, slug, match_title, filename))
                # PDFs without drawn table rules need a positional-word fallback.
                page_after = sum(len(v) for v in out.values())
                if page_after == page_before and slug:
                    merge_records(out, pdf_word_table_records(page, slug, filename))
    except Exception as e:
        print(f"WARNING: PDF parse failed for {filename}: {e}")
    return out, dates


def dedupe_and_validate(rows: List[dict]) -> Tuple[List[dict], float]:
    # Keep the last record for a duplicate symbol/name; repeated headers across PDF pages are common.
    dedup: Dict[str, dict] = {}
    for r in rows:
        name = clean_company(r.get("name"))
        if not name:
            continue
        symbol = norm_space(r.get("symbol"))
        key = compact(symbol) or compact(name)
        if not key:
            continue
        rec = {
            "name": name,
            "symbol": symbol,
            "sector": norm_space(r.get("sector")),
            "weight": round(float(r.get("weight") or 0), 6),
        }
        if rec["weight"] > 0:
            dedup[key] = rec
    cleaned = sorted(dedup.values(), key=lambda x: x["weight"], reverse=True)
    total = sum(r["weight"] for r in cleaned)
    # Some data sources encode 0-1 fractions. Convert if the whole portfolio is around 1.
    if cleaned and total <= 1.5:
        for r in cleaned:
            r["weight"] = round(r["weight"] * 100.0, 6)
        total = sum(r["weight"] for r in cleaned)
    return cleaned, total


def sector_rows(holdings: List[dict]) -> List[dict]:
    buckets: Dict[str, dict] = {}
    for h in holdings:
        sec = norm_space(h.get("sector"))
        if not sec:
            continue
        b = buckets.setdefault(sec, {"name": sec, "weight": 0.0, "count": 0})
        b["weight"] += float(h.get("weight") or 0)
        b["count"] += 1
    rows = list(buckets.values())
    for r in rows:
        r["weight"] = round(r["weight"], 6)
    return sorted(rows, key=lambda x: x["weight"], reverse=True)


def parse_report_zip(content: bytes, catalog: dict, report_label: str, source_url: str) -> Tuple[Dict[str, dict], dict]:
    match_title = title_matcher(catalog)
    all_rows: Dict[str, List[dict]] = {}
    asof_by_slug: Dict[str, str] = {}
    files_seen = []
    with zipfile.ZipFile(io.BytesIO(content)) as z:
        for info in z.infolist():
            if info.is_dir():
                continue
            name = info.filename
            suffix = Path(name).suffix.lower()
            if suffix not in {".csv", ".txt", ".xls", ".xlsx", ".pdf"}:
                continue
            files_seen.append(name)
            data = z.read(info)
            try:
                if suffix in {".csv", ".txt"}:
                    merge_records(all_rows, parse_csv_bytes(data, name, match_title))
                elif suffix in {".xls", ".xlsx"}:
                    merge_records(all_rows, parse_excel_bytes(data, name, match_title))
                elif suffix == ".pdf":
                    parsed, dates = parse_pdf_bytes(data, name, match_title)
                    merge_records(all_rows, parsed)
                    asof_by_slug.update(dates)
            except Exception as e:
                print(f"WARNING: parse failed for {name}: {e}")

    # Fallback report month date if exact date is not printed in the underlying file.
    m = re.search(r"([A-Z][a-z]{2})\s+(20\d{2})", report_label)
    fallback_asof = None
    if m:
        try:
            month = MONTH_ABBR.index(m.group(1)) + 1
            year = int(m.group(2))
            fallback_asof = date(year, month, calendar.monthrange(year, month)[1]).isoformat()
        except Exception:
            pass

    results: Dict[str, dict] = {}
    catalog_by_slug = catalog_index_map(catalog)
    for slug, item in catalog_by_slug.items():
        holdings, coverage = dedupe_and_validate(all_rows.get(slug, []))
        if not holdings:
            continue
        completeness = "full" if 97 <= coverage <= 103 else "partial"
        top10 = sum(h["weight"] for h in holdings[:10])
        data = {
            "slug": slug,
            "index_name": item.get("name"),
            "as_of": asof_by_slug.get(slug) or fallback_asof,
            "report_month": report_label,
            "source": "NSE Indices — Indices Market Capitalisation & Weightage",
            "source_url": source_url,
            "source_type": "official_nse_indices_monthly_report",
            "completeness": completeness,
            "stock_count": len(holdings),
            "weight_coverage": round(coverage, 4),
            "top10_weight": round(top10, 4),
            "coverage_note": (
                "Official NSE Indices monthly constituent-weight report."
                if completeness == "full"
                else f"Official report parsed, but extracted weights cover {coverage:.1f}% of the index. Treat this view as partial until a fuller table is available."
            ),
            "holdings": holdings,
            "sectors": sector_rows(holdings),
        }
        results[slug] = data

    diagnostics = {
        "report_label": report_label,
        "source_url": source_url,
        "files_seen": files_seen,
        "indices_parsed": len(results),
        "catalog_indices": len(catalog_by_slug),
        "missing_slugs": [s for s in catalog_by_slug if s not in results],
    }
    return results, diagnostics


def purge_demo_files(comp_dir: Path):
    for p in comp_dir.glob("*.json"):
        if p.name in {"manifest.json", "status.json"}:
            continue
        d = safe_json(p)
        if d and d.get("source_type") == "sample":
            print(f"Removing old demo composition file: {p.name}")
            p.unlink(missing_ok=True)


def build_manifest(catalog: dict, comp_dir: Path, parsed: Dict[str, dict], source_url: Optional[str], report_label: Optional[str], keep_existing: bool) -> dict:
    entries = {}
    for item in catalog.get("items", []):
        slug = item["slug"]
        path = comp_dir / f"{slug}.json"
        data = parsed.get(slug)
        if data:
            atomic_json(path, data)
        elif keep_existing and path.exists():
            old = safe_json(path)
            if old and str(old.get("source_type", "")).startswith("official"):
                data = old
        if data:
            entries[slug] = {
                "file": f"{slug}.json",
                "name": item.get("name"),
                "available": True,
                "as_of": data.get("as_of"),
                "completeness": data.get("completeness"),
                "stock_count": data.get("stock_count"),
                "weight_coverage": data.get("weight_coverage"),
                "top10_weight": data.get("top10_weight"),
                "source_type": data.get("source_type"),
            }
        else:
            entries[slug] = {
                "name": item.get("name"),
                "available": False,
                "reason": "Official constituent-weight data not parsed from the current monthly report.",
            }
    return {
        "version": 2,
        "generated_at": datetime.now(timezone.utc).replace(microsecond=0).isoformat().replace("+00:00", "Z"),
        "report_month": report_label,
        "source_url": source_url,
        "indices": entries,
    }


def autosize(ws, max_width: int = 42):
    for col in range(1, ws.max_column + 1):
        letter = get_column_letter(col)
        width = 0
        for row in range(1, min(ws.max_row, 600) + 1):
            v = ws.cell(row, col).value
            if v is not None:
                width = max(width, len(str(v)))
        ws.column_dimensions[letter].width = min(max(10, width + 2), max_width)


def update_workbook_composition(book_path: Path, data: Optional[dict]):
    try:
        wb = load_workbook(book_path)
    except Exception as e:
        print(f"WARNING: unable to open workbook {book_path}: {e}")
        return
    if "Composition" in wb.sheetnames:
        del wb["Composition"]
    ws = wb.create_sheet("Composition")
    ws.sheet_view.showGridLines = False
    dark = PatternFill("solid", fgColor="071A2D")
    header = PatternFill("solid", fgColor="0F2742")
    white = Font(color="FFFFFF", bold=True)
    muted = Font(color="666666")
    green = Font(color="008000")

    title = (data or {}).get("index_name") or book_path.stem.replace("-", " ").title()
    ws.merge_cells("A1:E1")
    ws["A1"] = f"{title} Composition"
    ws["A1"].fill = dark
    ws["A1"].font = Font(color="FFFFFF", bold=True, size=15)
    if not data:
        ws["A3"] = "Official composition data is not currently available for this index."
        ws["A3"].font = muted
        autosize(ws)
        wb.save(book_path)
        return

    ws["A2"] = "As of"
    ws["B2"] = data.get("as_of") or data.get("report_month")
    ws["D2"] = "Weight coverage"
    ws["E2"] = float(data.get("weight_coverage") or 0) / 100
    ws["E2"].number_format = "0.0%"
    ws["A3"] = "Source"
    ws["B3"] = data.get("source")
    ws["B3"].hyperlink = data.get("source_url")
    ws["B3"].style = "Hyperlink"
    ws["D3"] = "Top 10 weight"
    ws["E3"] = float(data.get("top10_weight") or 0) / 100
    ws["E3"].number_format = "0.0%"
    for c in ("A2", "D2", "A3", "D3"):
        ws[c].font = muted
    for c in ("B2", "E2", "E3"):
        ws[c].font = green

    headers = ["Rank", "Company", "Symbol", "Sector / Industry", "Weight"]
    for i, h in enumerate(headers, 1):
        cell = ws.cell(5, i, h)
        cell.fill = header
        cell.font = white
        cell.alignment = Alignment(horizontal="center")
    for r, h in enumerate(data.get("holdings", []), start=6):
        ws.cell(r, 1, r - 5)
        ws.cell(r, 2, h.get("name"))
        ws.cell(r, 3, h.get("symbol"))
        ws.cell(r, 4, h.get("sector"))
        ws.cell(r, 5, float(h.get("weight") or 0) / 100)
        ws.cell(r, 5).number_format = "0.00%"
    ws.freeze_panes = "A6"
    autosize(ws)

    # Put Composition before Sources & Methodology if that sheet exists.
    if "Sources & Methodology" in wb.sheetnames:
        comp = wb["Composition"]
        src = wb["Sources & Methodology"]
        wb._sheets.remove(comp)
        idx = wb._sheets.index(src)
        wb._sheets.insert(idx, comp)
    wb.save(book_path)


def update_all_workbooks(web_dir: Path, catalog: dict, manifest: dict):
    out_dir = web_dir / "downloads" / "indices"
    if not out_dir.exists():
        return
    comp_dir = web_dir / "data" / "composition"
    for item in catalog.get("items", []):
        slug = item["slug"]
        book = out_dir / f"{slug}.xlsx"
        if not book.exists():
            continue
        ent = manifest.get("indices", {}).get(slug, {})
        data = safe_json(comp_dir / ent.get("file", "")) if ent.get("file") else None
        update_workbook_composition(book, data)


def parse_args():
    p = argparse.ArgumentParser()
    p.add_argument("--web-dir", default="docs")
    p.add_argument("--max-months", type=int, default=8)
    p.add_argument("--strict", action="store_true", help="Exit non-zero if the official report cannot be refreshed")
    p.add_argument("--report-zip", help="Optional local report ZIP for testing/manual import")
    p.add_argument("--no-workbooks", action="store_true")
    return p.parse_args()


def main() -> int:
    args = parse_args()
    web_dir = Path(args.web_dir)
    catalog_path = web_dir / "data" / "multi_index_catalog.json"
    if not catalog_path.exists():
        print(f"Composition: catalog not found at {catalog_path}; run multi_index_tracker.py first.")
        return 1 if args.strict else 0
    catalog = json.loads(catalog_path.read_text(encoding="utf-8"))
    comp_dir = web_dir / "data" / "composition"
    comp_dir.mkdir(parents=True, exist_ok=True)
    purge_demo_files(comp_dir)

    existing_manifest = safe_json(comp_dir / "manifest.json") or {"indices": {}}
    parsed = {}
    diagnostics = {"status": "not_refreshed"}
    source_url = None
    report_label = None
    try:
        if args.report_zip:
            local = Path(args.report_zip)
            content = local.read_bytes()
            if not is_zip_bytes(content):
                raise RuntimeError(f"{local} is not a valid ZIP file")
            source_url = f"file://{local.resolve()}"
            # Best-effort report label from filename, otherwise current month.
            m = re.search(r"([A-Z][a-z]{2})(20\d{2})", local.name)
            report_label = f"{m.group(1)} {m.group(2)}" if m else f"{MONTH_ABBR[date.today().month-1]} {date.today().year}"
            print(f"Composition: importing local report {local} ({len(content):,} bytes)")
        else:
            content, source_url, report_label = fetch_report_zip(date.today(), args.max_months)
            print(f"Composition: downloaded {report_label} report from {source_url} ({len(content):,} bytes)")
        parsed, diagnostics = parse_report_zip(content, catalog, report_label, source_url)
        diagnostics["status"] = "ok"
        print(f"Composition: parsed {len(parsed)}/{len(catalog.get('items', []))} tracked indices")
    except Exception as e:
        diagnostics = {"status": "refresh_failed", "error": str(e)}
        print(f"WARNING: composition refresh failed: {e}")
        if args.strict:
            return 2

    manifest = build_manifest(
        catalog,
        comp_dir,
        parsed,
        source_url or existing_manifest.get("source_url"),
        report_label or existing_manifest.get("report_month"),
        keep_existing=True,
    )
    atomic_json(comp_dir / "manifest.json", manifest)
    status = {
        **diagnostics,
        "generated_at": datetime.now(timezone.utc).replace(microsecond=0).isoformat().replace("+00:00", "Z"),
        "available_indices": sum(1 for x in manifest["indices"].values() if x.get("available")),
        "tracked_indices": len(catalog.get("items", [])),
    }
    atomic_json(comp_dir / "status.json", status)

    if not args.no_workbooks:
        update_all_workbooks(web_dir, catalog, manifest)
        print("Composition: Excel Composition sheets refreshed.")

    print(f"Composition: {status['available_indices']}/{status['tracked_indices']} indices available in manifest.")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
